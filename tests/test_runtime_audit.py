import io
import os
import re
import tempfile

os.environ.setdefault('REVIEW_DATA_DIR', tempfile.mkdtemp(prefix='runtime-audit-'))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app, startup
from backend.db import DATA, Issue, Member, Project, Run, Session, Task, User


def _login(client):
    text = (DATA / 'bootstrap-admin.txt').read_text(encoding='utf-8')
    password = re.findall(r'[A-Za-z0-9_-]{20,}', text)[-1]
    result = client.post('/api/login', json={'username': 'admin', 'password': password})
    assert result.status_code == 200, result.text
    return {'X-CSRF-Token': result.json()['csrf']}


def _risk_fixture():
    startup()
    with Session() as db:
        user = db.query(User).filter_by(username='admin').first()
        project = Project(name='runtime risk', owner_id=user.id, config={})
        db.add(project)
        db.flush()
        db.add(Member(project_id=project.id, user_id=user.id, role='owner'))
        run = Run(project_id=project.id, request_key='runtime-risk', mode='strict', modules=['high_risk'], snapshot={'document_ids': {}}, status='done')
        db.add(run)
        db.flush()
        output = {
            'source_rows': [
                {'row_id': 's1', 'chapter': 'CH03', 'table_location': 'report:T3', 'row_number': 2, 'kind': 'problem', 'parse_status': 'parsed', 'object': 'server-1', 'description': 'weak password', 'analysis': 'high risk', 'grade': 'high', 'raw_values': {'description': 'weak password'}},
                {'row_id': 's2', 'chapter': 'CH05', 'table_location': 'report:T5', 'row_number': 3, 'kind': 'risk', 'parse_status': 'parsed', 'object': 'server-2', 'description': 'grade mismatch', 'analysis': 'high risk', 'grade': 'low', 'raw_values': {'description': 'grade mismatch'}},
            ],
            'guide_rows': [{'key': 'g1', 'scenario': 'weak password', 'source': 'guide:1'}],
            'links': [{'link_id': 's1:g1', 'source_row_id': 's1', 'guide_key': 'g1', 'score': 1, 'applicability': 'pending', 'status': 'candidate', 'machine_conclusion': None}],
            'candidates': [{'row_id': 's1', 'description': 'weak password', 'matches': [{'key': 'g1'}], 'status': 'candidate'}],
            'unmatched': [{'row_id': 's2', 'source_row_id': 's2', 'source': 'report:T5', 'reason': 'manual review'}],
            'conflicts': [{'type': 'grade_conflict', 'row_id': 's2', 'expected': 'high', 'actual': 'low', 'source': 'report:T5'}],
            'applicability': {'g1': 'pending'},
            'linked_rows': [],
        }
        task = Task(project_id=project.id, run_id=run.id, kind='risk', label='risk', payload={}, status='done', output=output)
        db.add(task)
        db.flush()
        db.add(Issue(project_id=project.id, run_id=run.id, task_id=task.id, chapter='CH05', category='risk_consistency', title='conflict', description='grade mismatch', suggestion='review', object_name='server-2', evidence=[], machine={'row_id': 's2'}))
        db.commit()
        return project.id, run.id


def test_risk_api_review_and_export_preserve_four_way_rows():
    pid, rid = _risk_fixture()
    with TestClient(app) as client:
        headers = _login(client)
        response = client.get(f'/api/projects/{pid}/runs/{rid}/risk')
        assert response.status_code == 200
        assert {row['row_id'] for row in response.json()['source_rows']} == {'s1', 's2'}
        reviewed = client.patch(f'/api/projects/{pid}/runs/{rid}/risk/candidate:s1', headers=headers,
                                json={'version': 0, 'status': 'confirmed', 'conclusion': '人工确认', 'reason': '证据一致'})
        assert reviewed.status_code == 200, reviewed.text
        body = client.get(f'/api/projects/{pid}/runs/{rid}/risk').json()
        assert body['candidates'][0]['review']['conclusion'] == '人工确认'
        assert body['candidates'][0]['review']['reason'] == '证据一致'
        exported = client.get(f'/api/projects/{pid}/runs/{rid}/risk/export')
        assert exported.status_code == 200, exported.text
        book = load_workbook(io.BytesIO(exported.content), read_only=True)
        assert {'候选关联', '未匹配来源', '冲突与适用性', '四方核对证据'} <= set(book.sheetnames)
        candidate_values = list(book['候选关联'].values)
        assert any('人工确认' in row for row in candidate_values)


def test_risk_api_and_export_include_legacy_thirteen_column_summary():
    pid, rid = _risk_fixture()
    with Session() as db:
        task = db.query(Task).filter_by(run_id=rid, kind='risk').one()
        output = dict(task.output)
        source_rows = list(output['source_rows'])
        source_rows[0] = {**source_rows[0], 'values': {
            '问题编号': 'P-1', '安全问题': 'weak password', '测评项': 'a）应实施身份鉴别'}}
        source_rows.append({**source_rows[0], 'row_id': 's3', 'description': 'another password issue',
                            'values': {'问题编号': 'P-2', '安全问题': 'another password issue',
                                       '测评项': '应实施身份鉴别'}})
        output['source_rows'] = source_rows
        output['guide_rows'] = [{'key': 'g1', 'clause': '6.1', 'requirement': '应实施身份鉴别',
                                 'scope': '三级', 'scenario': '弱口令情形',
                                 'mitigation': '改进认证', 'evaluation': '参考评价', 'major': '高风险项'}]
        task.output = output
        db.commit()
    with TestClient(app) as client:
        _login(client)
        result = client.get(f'/api/projects/{pid}/runs/{rid}/risk').json()
        assert len(result['legacy_rows']) == 2
        assert result['legacy_rows'][0]['values']['问题编号'] == 'P-1'
        assert result['legacy_rows'][1]['values']['指引-场景/问题描述'] == '弱口令情形'
        book = load_workbook(io.BytesIO(client.get(f'/api/projects/{pid}/runs/{rid}/risk/export').content),
                             read_only=False)
        sheet = book['高风险核查汇总']
        assert sheet.max_column == 13
        assert sheet.max_row == 3
        assert [cell.value for cell in sheet[1]][:3] == ['序号', '问题编号', '报告安全问题描述']
        assert sheet['B2'].value == 'P-1'
        assert sheet['B3'].value == 'P-2'
        assert {str(area) for area in sheet.merged_cells.ranges} == {'K2:K3', 'L2:L3', 'M2:M3'}
        assert all(sheet.cell(3, col).value is None for col in (11, 12, 13))


def test_issue_listing_and_export_can_be_bound_to_run_and_module():
    pid, rid = _risk_fixture()
    with Session() as db:
        other = Run(project_id=pid, request_key='runtime-other', mode='strict', modules=['appendix_d'], snapshot={'document_ids': {}}, status='done')
        db.add(other)
        db.flush()
        db.add(Issue(project_id=pid, run_id=other.id, chapter='APP_D', category='typo', title='other', description='other', suggestion='', object_name='', evidence=[], machine={'status': 'pending'}))
        current_task = db.query(Task).filter_by(run_id=rid, kind='risk').first()
        db.add(Issue(project_id=pid, run_id=rid, task_id=current_task.id, chapter='CH05', category='risk_consistency', title='confirmed current', description='current', suggestion='', object_name='', evidence=[], machine={}, status='confirmed'))
        db.commit()
        other_id = other.id
    with TestClient(app) as client:
        _login(client)
        current = client.get(f'/api/projects/{pid}/issues?run_id={rid}&module=high_risk')
        assert current.status_code == 200
        assert current.json() and all(row['run_id'] == rid for row in current.json())
        assert client.get(f'/api/projects/{pid}/issues?run_id={rid}&module=appendix_d').json() == []
        assert client.get(f'/api/projects/{pid}/issues?run_id={other_id}').status_code == 200
        empty = client.get(f'/api/projects/{pid}/issues?run_id=')
        assert empty.status_code == 200 and empty.json() == []
        exported = client.get(f'/api/projects/{pid}/export?format=xlsx&run_id={rid}&module=high_risk')
        assert exported.status_code == 200
        sheet = load_workbook(io.BytesIO(exported.content), read_only=True).active
        assert sheet.max_row == 2 and sheet.cell(2, 3).value == 'confirmed current'
        filtered = client.get(f'/api/projects/{pid}/export?format=xlsx&run_id={rid}&module=appendix_d')
        assert filtered.status_code == 200
        assert load_workbook(io.BytesIO(filtered.content), read_only=True).active.max_row == 1
