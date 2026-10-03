import io
import os
import tempfile

os.environ.setdefault('REVIEW_DATA_DIR', tempfile.mkdtemp(prefix='appendix-api-'))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app, startup
from backend.db import DATA, Session, Project, Run, Task, Issue, User, Member


def test_appendix_view_is_one_row_per_record_and_confirmed_export_filters():
    startup()
    with Session() as db:
        user=db.query(User).filter_by(username='admin').first()
        project=Project(name='synthetic appendix',owner_id=user.id,config={})
        db.add(project);db.flush();db.add(Member(project_id=project.id,user_id=user.id,role='owner'))
        run=Run(project_id=project.id,request_key='appendix-api',mode='strict',modules=['appendix_d'],snapshot={'document_ids':{}},status='done')
        db.add(run);db.flush()
        task=Task(project_id=project.id,run_id=run.id,kind='record',label='r1',payload={'record_id':'r1'},status='done',output={'record':{'id':'r1','chapter_number':'附录D.1','domain':'安全通信网络','layer_order':1,'object':'对象A','object_order':1,'source_order':1,'control':'控制点','requirement':'测评项','text':'=SUM(1,1)','verdict':'符合'},'requirement':{'points':[{'text':'核查点'}]}})
        db.add(task);db.flush()
        db.add(Issue(project_id=project.id,run_id=run.id,task_id=task.id,chapter='APP_D',category='verdict',title='J',description='矛盾',suggestion='复核',object_name='对象A',evidence=[],machine={'record_id':'r1','column':'J'},status='confirmed'))
        db.add(Issue(project_id=project.id,run_id=run.id,task_id=task.id,chapter='APP_D',category='typo',title='M',description='错字',suggestion='修正',object_name='对象A',evidence=[],machine={'record_id':'r1','column':'M'},status='pending'))
        db.commit();pid,rid=project.id,run.id
        password=(DATA/'bootstrap-admin.txt').read_text(encoding='utf-8').split('初始密码：')[1].splitlines()[0]
    with TestClient(app) as client:
        csrf=client.post('/api/login',json={'username':'admin','password':password}).json()['csrf'];headers={'X-CSRF-Token':csrf}
        view=client.get(f'/api/projects/{pid}/runs/{rid}/appendix?view=issues').json()
        assert len(view['rows'])==1 and len(view['rows'][0]['issues']['verdict'])==1 and len(view['rows'][0]['issues']['typo'])==1
        exported=client.get(f'/api/projects/{pid}/runs/{rid}/appendix/export?view=confirmed')
        book=load_workbook(io.BytesIO(exported.content));sheet=book['附录D结果记录核查']
        assert sheet.max_row==2 and sheet['G2'].value.startswith("'=SUM")
