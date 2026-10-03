import io
import os
import tempfile

os.environ['REVIEW_DATA_DIR'] = tempfile.mkdtemp(prefix='consistency-api-')

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest


def load_backend():
    # Delay the DB import until collection has finished; test_core sets its own test data dir.
    global app, password_hash, Base, Document, Member, Project, Run, Session, Task, User, engine
    from backend.app import app, password_hash
    from backend.db import Base, Document, Member, Project, Run, Session, Task, User, engine


@pytest.fixture(autouse=True)
def clean_tables():
    yield
    load_backend()
    Base.metadata.drop_all(engine)


def setup_data(rows=None, diagnostics=None):
    load_backend()
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with Session() as db:
        owner = User(username='owner', display_name='Owner', password_hash=password_hash('password123456'))
        other = User(username='other', display_name='Other', password_hash=password_hash('password123456'))
        db.add_all([owner, other]); db.flush()
        project = Project(name='Synthetic', owner_id=owner.id, config={})
        db.add(project); db.flush()
        db.add_all([Member(project_id=project.id, user_id=u.id) for u in (owner, other)])
        run = Run(project_id=project.id, request_key='one', mode='lenient', modules=['assets_full'],
                  snapshot={'document_ids': {}, 'project_config': {}}, status='done')
        db.add(run); db.flush()
        task = Task(project_id=project.id, run_id=run.id, kind='assets', label='Synthetic',
                    payload={'scope': 'full'}, status='done', output={
                        'scope': 'full', 'mapping_version': 'test-v1', 'rows': rows or [],
                        'issues': [], 'diagnostics': diagnostics or [], 'completeness': not diagnostics})
        db.add(task); db.commit()
        return project.id, run.id, task.id


def client_for(username='owner'):
    client = TestClient(app)
    response = client.post('/api/login', json={'username': username, 'password': 'password123456'})
    assert response.status_code == 200, response.text
    return client, {'X-CSRF-Token': response.json()['csrf']}


def sample_rows():
    return [
        {'id': 'r1', 'scope': 'full', 'category_id': 'c01', 'category': '机房',
         'entity_key': '001', 'status': 'different', 'has_difference': True, 'incomplete': False,
         'fields': {'f1': {'field_id': 'f1', 'label': '位置', 'status': 'different',
                           'different_sources': ['survey', 'plan'],
                           'issue_codes': ['different_value'], 'sources': {
                               'survey': {'raw_value': '=1+1', 'status': 'value', 'source': {'table_id': 't1'}},
                               'plan': {'raw_value': '+cmd', 'status': 'value', 'source': {'table_id': 't2'}},
                               'report': {'raw_value': '00123', 'status': 'value', 'source': {'table_id': 't3'}}}}}},
        {'id': 'r2', 'scope': 'full', 'category_id': 'c02', 'category': '设备',
         'entity_key': 'second', 'status': 'consistent', 'has_difference': False,
         'incomplete': False, 'fields': {}}
    ]


def test_new_upload_keeps_old_run_source_versions():
    pid, rid, tid = setup_data()
    with Session() as db:
        old = Document(project_id=pid, role='report', filename='old.docx', sha256='a'*64,
                       storage_key='old', version=1, status='parsed', parsed={'source_tables': []})
        db.add(old); db.flush()
        run = db.get(Run, rid); run.snapshot = {'document_ids': {'report': old.id}, 'project_config': {}}
        db.commit()
        newer = Document(project_id=pid, role='report', filename='new.docx', sha256='b'*64,
                         storage_key='new', version=2, status='parsed', parsed={'source_tables': []})
        db.add(newer); db.commit()
        assert db.get(Run, rid).snapshot['document_ids']['report'] == old.id
    client, _ = client_for()
    response = client.get(f'/api/projects/{pid}/runs/{rid}/consistency?scope=full')
    assert response.status_code == 200, response.text
    assert response.json()['rows'] == []


def test_consistency_page_and_export_share_row_ids_and_order():
    pid, rid, _ = setup_data(sample_rows())
    client, _ = client_for('other')
    page = client.get(f'/api/projects/{pid}/runs/{rid}/consistency?scope=full&category=c01')
    assert page.status_code == 200, page.text
    assert [r['id'] for r in page.json()['rows']] == ['r1']
    response = client.get(f'/api/projects/{pid}/runs/{rid}/consistency/export?scope=full&category=c01')
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content))['逐对象逐字段结果']
    assert [sheet.cell(i, 1).value for i in range(3, sheet.max_row + 1)] == ['r1', 'r1', 'r1']


def test_missing_source_table_disables_all_consistent():
    pid, rid, _ = setup_data(diagnostics=[{'position_id': 'c01:report:full', 'status': 'missing'}])
    client, _ = client_for()
    page = client.get(f'/api/projects/{pid}/runs/{rid}/consistency?scope=full')
    assert page.status_code == 200, page.text
    assert page.json()['completeness'] is False
    assert page.json()['diagnostics'][0]['status'] == 'missing'


def test_export_has_two_review_sheets_and_filter_label():
    pid, rid, _ = setup_data(sample_rows())
    client, _ = client_for()
    response = client.get(f'/api/projects/{pid}/runs/{rid}/consistency/export?scope=full&status=different')
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ['三文档横向汇总', '逐对象逐字段结果', '异常明细']
    assert 'different' in str(book.worksheets[0]['A1'].value)


def test_export_escapes_formula_prefixes_and_keeps_leading_zero():
    pid, rid, _ = setup_data(sample_rows())
    client, _ = client_for()
    book = load_workbook(io.BytesIO(client.get(f'/api/projects/{pid}/runs/{rid}/consistency/export').content))
    values = [cell.value for sheet in book for row in sheet for cell in row]
    assert "'=1+1" in values and "'+cmd" in values
    assert '00123' in values
    assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)


def test_export_colors_only_sources_that_differ_from_report():
    pid, rid, _ = setup_data(sample_rows())
    client, _ = client_for()
    sheet = load_workbook(io.BytesIO(client.get(f'/api/projects/{pid}/runs/{rid}/consistency/export').content))['逐对象逐字段结果']
    assert [sheet.cell(row, 9).value for row in (3, 4, 5)] == ['survey', 'plan', 'report']
    assert sheet.cell(3, 10).fill.fgColor.rgb == '00FFF2CC'
    assert sheet.cell(4, 10).fill.fgColor.rgb == '00FFF2CC'
    assert sheet.cell(5, 10).fill.patternType is None


def test_export_wide_sheet_puts_same_field_sources_on_one_row():
    pid, rid, _ = setup_data(sample_rows())
    client, _ = client_for()
    book = load_workbook(io.BytesIO(client.get(f'/api/projects/{pid}/runs/{rid}/consistency/export').content))
    sheet = book['三文档横向汇总']
    row = next(row for row in sheet if any(cell.value == "'=1+1" for cell in row))
    values = [cell.value for cell in row]
    offset = values.index("'=1+1")
    assert values[offset:offset + 3] == ["'=1+1", "'+cmd", '00123']
    assert row[offset].fill.fgColor.rgb == '00FFF2CC'
    assert row[offset + 1].fill.fgColor.rgb == '00FFF2CC'
    assert row[offset + 2].fill.patternType is None


def test_review_conflict_returns_409():
    pid, rid, tid = setup_data(sample_rows())
    client, headers = client_for('other')
    url = f'/api/projects/{pid}/runs/{rid}/consistency/r1'
    first = client.patch(url, json={'version': 0, 'status': 'confirmed', 'note': 'Checked'}, headers=headers)
    assert first.status_code == 200, first.text
    assert first.json()['version'] == 1
    assert client.patch(url, json={'version': 0, 'status': 'rejected'}, headers=headers).status_code == 409
    with Session() as db:
        assert db.get(Task, tid).output['rows'] == sample_rows()


def test_precheck_uses_source_tables_instead_of_legacy_asset_rows():
    pid, _, _ = setup_data()
    with Session() as db:
        for role in ('survey', 'plan', 'report'):
            db.add(Document(project_id=pid, role=role, filename=f'{role}.docx',
                            sha256=role, storage_key=role, version=1, status='parsed',
                            parsed={'source_tables': [], 'assets': [], 'counts': {}}))
        db.commit()
    client, headers = client_for()
    response = client.post(f'/api/projects/{pid}/precheck',
                           json={'modules': ['assets_full', 'assets_sample']}, headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()['blockers'] == []


def test_legacy_asset_run_requires_new_consistency_run():
    pid, rid, tid = setup_data()
    with Session() as db:
        task = db.get(Task, tid)
        task.output = {'scope': 'full', 'rows': [{'name': 'old asset row'}]}
        db.commit()
    client, headers = client_for()
    response = client.get(f'/api/projects/{pid}/runs/{rid}/consistency')
    assert response.status_code == 409
    assert '旧运行' in response.json()['detail']
    reviewed = client.patch(f'/api/projects/{pid}/runs/{rid}/consistency/old-row',
                            json={'version': 0, 'status': 'confirmed'}, headers=headers)
    assert reviewed.status_code == 404


def test_queued_run_uses_its_document_and_creation_snapshots(monkeypatch):
    pid, rid, tid = setup_data()
    with Session() as db:
        old = Document(project_id=pid, role='report', filename='old.docx', sha256='a'*64,
                       storage_key='old', version=1, status='parsed', parsed={'source_tables': []})
        db.add(old); db.flush()
        run = db.get(Run, rid)
        run.snapshot = {'document_ids': {'report': old.id},
                        'document_versions': {'report': {'id': old.id, 'version': 1}},
                        'created_at': 123.0, 'algorithm_version': 'consistency-v1'}
        new = Document(project_id=pid, role='report', filename='new.docx', sha256='b'*64,
                       storage_key='new', version=2, status='parsed', parsed={'source_tables': []})
        db.add(new); db.commit()
        seen = []
        monkeypatch.setattr('backend.app.compare_consistency',
                            lambda docs, scope: (seen.append(docs['report']['id']) or
                                                 {'mapping_version': 'test-v1', 'rows': []}))
        from backend.app import perform_task
        result = perform_task(db, db.get(Task, tid))['output']
        assert seen == [old.id]
        assert result['document_versions']['report']['version'] == 1
        assert result['created_at'] == 123.0
        assert len(result['categories']) == 12


def test_difference_filter_keeps_rows_that_are_also_incomplete():
    rows = sample_rows()
    rows[0]['status'] = 'missing_object'
    rows[0]['incomplete'] = True
    pid, rid, _ = setup_data(rows)
    client, _ = client_for()
    base = f'/api/projects/{pid}/runs/{rid}/consistency'
    page = client.get(base + '?status=different')
    assert page.status_code == 200, page.text
    assert [row['id'] for row in page.json()['rows']] == ['r1']
    export = client.get(base + '/export?status=different')
    sheet = load_workbook(io.BytesIO(export.content))['逐对象逐字段结果']
    assert {sheet.cell(i, 1).value for i in range(3, sheet.max_row + 1)} == {'r1'}


def test_parse_failed_filter_finds_failed_source_cell():
    rows = sample_rows()
    rows[0]['status'] = 'incomplete'
    rows[0]['fields']['f1']['sources']['report']['status'] = 'parse_failed'
    pid, rid, _ = setup_data(rows)
    client, _ = client_for()
    page = client.get(f'/api/projects/{pid}/runs/{rid}/consistency?status=parse_failed')
    assert page.status_code == 200, page.text
    assert [row['id'] for row in page.json()['rows']] == ['r1']
