import io

import pytest
from openpyxl import Workbook

from backend.catalog import import_workbook


def workbook_with_group_heading(include_requirement=True):
    book = Workbook()
    sheet = book.active
    sheet.title = '安全通信网络'
    sheet.append(['序号', '控制点', '属性标识SAG', '测评项', '点数', '核查点', '判定规则'])
    sheet.merge_cells('A2:G2')
    sheet['A2'] = '安全测评通用要求'
    if include_requirement:
        sheet.append([1, '网络架构', 'G', '应保证设备容量', 1, '应核查设备容量', '满足为符合'])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_merged_group_heading_is_not_a_requirement_or_warning():
    result = import_workbook(workbook_with_group_heading(), '安全通用要求_三级_核查点梳理统计表格.xlsx')

    assert [item['text'] for item in result['requirements']] == ['应保证设备容量']
    assert result['warnings'] == []


def test_workbook_with_only_group_heading_has_no_valid_requirement():
    with pytest.raises(ValueError, match='未找到有效测评项'):
        import_workbook(workbook_with_group_heading(False), '安全通用要求_三级_核查点梳理统计表格.xlsx')


def test_high_risk_guide_imports_requirement_used_by_legacy_match():
    book = Workbook()
    sheet = book.active
    sheet.title = '扩充文档'
    sheet.append(['高风险条款号', '要求项/标准要求', '适用范围', '场景/问题描述',
                  '可能的缓解措施', '风险评价', '是否重大风险'])
    sheet.append(['6.1.1', '应对登录用户进行身份鉴别', '三级', '未进行身份鉴别的情形',
                  '增加认证', '风险评价参考', '高风险项'])
    sheet.append(['', '', '', '身份鉴别失败的另一情形', '', '', ''])
    stream = io.BytesIO()
    book.save(stream)

    result = import_workbook(stream.getvalue(), '高风险判定指引.xlsx')
    assert result['requirements'][0]['requirement'] == '应对登录用户进行身份鉴别'
    assert len(result['requirements']) == 2
    assert result['requirements'][1]['clause'] == '6.1.1'
    assert result['requirements'][1]['requirement'] == '应对登录用户进行身份鉴别'
    assert result['requirements'][1]['scenario'] == '身份鉴别失败的另一情形'
