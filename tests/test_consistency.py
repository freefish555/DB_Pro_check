from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.documents import parse_docx
from backend.consistency import compare_documents


def parsed(doc, role='plan'):
    stream = BytesIO()
    doc.save(stream)
    return parse_docx(stream.getvalue(), role)


def inventory(doc):
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['设备名称', 'IP 地址', '重要程度']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['交换机 A', 'https://host:443/path', '重要']):
        cell.text = value
    return table


def test_source_table_keeps_header_cell_and_section_path():
    doc = Document()
    doc.add_heading('2 测评对象', 1)
    doc.add_heading('2.4 全对象清单', 2)
    doc.add_paragraph('表 2-1 网络设备')
    inventory(doc)
    source = parsed(doc)['source_tables'][0]
    assert source['chapter_path'] == ['2 测评对象', '2.4 全对象清单']
    assert source['parent_titles'] == source['chapter_path']
    assert source['caption'] == '表 2-1 网络设备'
    assert source['headers'] == ['设备名称', 'IP 地址', '重要程度']
    cell = next(c for c in source['cells'] if (c['row'], c['col']) == (2, 2))
    assert cell['raw_label'] == 'IP 地址'
    assert cell['raw_value'] == 'https://host:443/path'
    assert cell['source_cell'] == 't1r2c2'
    assert source['table_id'] == 't1'
    assert source['parse_status'] == 'parsed'
    assert 'CH02' in source['location']


def test_plan_chapter_three_table_two_caption_is_sample_context():
    doc = Document()
    doc.add_heading('第3章 测评对象选择结果', 1)
    doc.add_heading('3.1 网络设备', 2)
    doc.add_paragraph('表 2-1 网络设备')
    inventory(doc)
    result = parsed(doc)
    assert result['source_tables'][0]['chapter_path'][0] == '第3章 测评对象选择结果'
    assert result['assets'][0]['scope'] == 'planned_sample'


def test_merged_grid_preserves_origin_and_does_not_fill_plain_blank():
    doc = Document()
    doc.add_heading('2 清单', 1)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '合并标题'
    table.cell(1, 0).merge(table.cell(2, 0)).text = '纵向值'
    table.cell(1, 1).text = '独立值'
    cells = {(c['row'], c['col']): c for c in parsed(doc)['source_tables'][0]['cells']}
    assert cells[1, 2]['source_cell'] == 't1r1c1'
    assert cells[1, 1]['grid_span'] == 2
    assert cells[2, 1]['v_merge'] == 'restart'
    assert cells[3, 1]['v_merge'] == 'continue'
    assert cells[3, 1]['source_cell'] == 't1r2c1'
    assert cells[3, 1]['raw_value'] == ''
    assert cells[3, 2]['source_cell'] == 't1r3c2'
    assert cells[3, 2]['raw_value'] == ''


def test_nested_table_is_diagnostic_not_asset():
    doc = Document()
    doc.add_heading('2 清单', 1)
    doc.add_paragraph('表 2-1 网络设备')
    table = inventory(doc)
    nested = table.cell(1, 0).add_table(rows=2, cols=3)
    for cell, value in zip(nested.rows[0].cells, ['设备名称', 'IP 地址', '重要程度']):
        cell.text = value
    nested.cell(1, 0).text = '嵌套设备'
    result = parsed(doc)
    assert result['assets'] == []
    assert result['source_tables'][0]['parse_status'] == 'nested_table'
    assert any('嵌套表' in diagnostic for diagnostic in result['diagnostics'])


def test_empty_table_keeps_unresolved_source_evidence():
    doc = Document()
    doc.add_table(rows=0, cols=3)
    result = parsed(doc)
    source = result['source_tables'][0]
    assert source['parse_status'] == 'unknown_section'
    assert source['chapter_path'] == []
    assert source['cells'] == []
    assert result['assets'] == []
    assert any('不可据此认定不存在资产' in diagnostic for diagnostic in result['diagnostics'])


def test_custom_outline_heading_sets_sample_chapter_despite_table_two_caption():
    doc = Document()
    style = doc.styles.add_style('Custom Chapter', WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '0')
    style.element.get_or_add_pPr().append(outline)
    doc.add_paragraph('第3章 测评对象选择结果', style=style)
    doc.add_paragraph('表 2-1 网络设备')
    inventory(doc)
    result = parsed(doc)
    assert result['source_tables'][0]['chapter_path'] == ['第3章 测评对象选择结果']
    assert result['sections'][0]['code'] == 'CH03'
    assert result['source_tables'][0]['location'] == 'CH03 · 表1'
    assert result['assets'][0]['scope'] == 'planned_sample'


def test_grid_before_keeps_header_labels_at_logical_columns():
    doc = Document()
    doc.add_heading('2 清单', 1)
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 1).text = '设备名称'
    table.cell(0, 2).text = 'IP 地址'
    table.cell(1, 1).text = '交换机'
    table.cell(1, 2).text = '10.0.0.1'
    row = table._tbl.tr_lst[0]
    row.remove(row.tc_lst[0])
    before = OxmlElement('w:gridBefore')
    before.set(qn('w:val'), '1')
    row.get_or_add_trPr().append(before)
    source = parsed(doc)['source_tables'][0]
    assert source['headers'] == ['', '设备名称', 'IP 地址']
    cells = {(c['row'], c['col']): c for c in source['cells']}
    assert cells[1, 2]['raw_label'] == '设备名称'
    assert cells[1, 3]['raw_label'] == 'IP 地址'
    assert cells[2, 1]['raw_label'] == ''
    assert cells[2, 2]['raw_label'] == '设备名称'
    assert cells[2, 3]['raw_label'] == 'IP 地址'
    assert cells[1, 2]['source_cell'] == 't1r1c2'


def room_source(values=None, scope='full'):
    doc = Document()
    doc.add_heading('2 测评对象' if scope == 'full' else '第3章 测评对象选择结果', 1)
    doc.add_heading('2.4 系统构成' if scope == 'full' else '3.1 测评对象选择结果', 2)
    doc.add_heading('物理机房', 3)
    doc.add_paragraph('表 2-99 物理机房')
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ['机房名称', '物理位置', '重要程度']):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, values or ['机房 A', '园区', '重要']):
        cell.text = text
    return parsed(doc)


def test_mapping_has_12_classes_65_fields_and_58_source_positions():
    from backend.consistency_mapping import MAPPING
    assert MAPPING['mapping_version']
    assert len(MAPPING['categories']) == 12
    assert sum(len(c['fields']) for c in MAPPING['categories']) == 65
    assert len(MAPPING['positions']) == 58
    assert 'case_body_table' not in str(MAPPING)
    assert 'case_physical_cell' not in str(MAPPING)
    assert {p['role'] for p in MAPPING['positions']} == {'survey', 'plan', 'report'}
    assert {p['scope'] for p in MAPPING['positions']} == {'full', 'sample'}


def test_locator_prefers_chapter_parent_and_headers_over_caption_number():
    from backend.consistency_mapping import locate_tables
    result = room_source(scope='sample')
    full = locate_tables(result, 'plan', 'full')['c01:plan:full']
    sample = locate_tables(result, 'scheme', 'selected')['c01:plan:sample']
    assert full['status'] == 'missing'
    assert sample['status'] == 'matched'
    assert sample['table_ids'] == ['t1']


def test_ambiguous_table_is_not_silently_selected():
    from backend.consistency_mapping import locate_tables, extract_cells
    result = room_source()
    duplicate = dict(result['source_tables'][0], table_id='t99')
    result['source_tables'].append(duplicate)
    match = locate_tables(result, 'plan', 'full')['c01:plan:full']
    assert match['status'] == 'ambiguous'
    assert match['table_ids'] == ['t1', 't99']
    assert all(c['status'] == 'ambiguous' for c in extract_cells(result, 'plan', 'full') if c['category_id'] == 'c01')


def test_stronger_header_signature_selects_one_table():
    from backend.consistency_mapping import locate_tables
    result = room_source()
    weaker = dict(result['source_tables'][0], table_id='t99',
                  cells=[dict(cell) for cell in result['source_tables'][0]['cells']])
    next(cell for cell in weaker['cells'] if cell['row'] == 1 and cell['col'] == 3)['raw_value'] = '其他列'
    result['source_tables'].append(weaker)
    match = locate_tables(result, 'plan', 'full')['c01:plan:full']
    assert match['status'] == 'matched'
    assert match['table_ids'] == ['t1']


def test_not_applicable_explanation_row_is_not_asset():
    from backend.consistency_mapping import extract_cells
    result = room_source(['本次测评不涉及', '', ''])
    cells = [c for c in extract_cells(result, 'plan', 'full') if c['category_id'] == 'c01']
    assert len(cells) == 1
    assert cells[0]['status'] == 'not_applicable'
    assert cells[0]['entity_name'] is None
    assert cells[0]['source']['row'] == 2


def test_merged_explanation_row_is_one_not_applicable_marker():
    from backend.consistency_mapping import extract_cells
    doc = Document()
    doc.add_heading('2 测评对象', 1)
    doc.add_heading('2.4 系统构成', 2)
    doc.add_heading('物理机房', 3)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['机房名称', '物理位置', '重要程度']):
        cell.text = value
    table.cell(1, 0).merge(table.cell(1, 2)).text = '本次测评不涉及物理机房'
    cells = [c for c in extract_cells(parsed(doc), 'plan', 'full') if c['category_id'] == 'c01']
    assert len(cells) == 1
    assert cells[0]['status'] == 'not_applicable'


def test_nonempty_row_without_name_is_parse_issue():
    from backend.consistency_mapping import extract_cells
    cells = [c for c in extract_cells(room_source(['', '园区', '重要']), 'plan', 'full') if c['category_id'] == 'c01']
    assert len(cells) == 1
    assert cells[0]['status'] == 'parse_failed'
    assert cells[0]['diagnostic'] == 'nonempty_row_without_name'
    assert cells[0]['source']['row'] == 2


def test_extraction_preserves_blank_value_and_cell_origin():
    from backend.consistency_mapping import extract_cells
    cells = [c for c in extract_cells(room_source(['机房 A', '', '重要']), 'plan', 'full') if c['category_id'] == 'c01']
    assert len(cells) == 3
    assert cells[1]['status'] == 'empty'
    assert cells[1]['raw_value'] == ''
    assert cells[1]['source']['source_cell'] == 't1r2c2'


def test_numeric_chapter_conflict_is_parse_failed():
    from backend.consistency_mapping import locate_tables
    result = room_source()
    result['source_tables'][0]['chapter_path'][0] = '4 测评方法'
    result['source_tables'][0]['parent_titles'][0] = '4 测评方法'
    match = locate_tables(result, 'plan', 'full')['c01:plan:full']
    assert match['status'] == 'parse_failed'
    assert 'chapter_conflict' in match['diagnostics']


def test_plan_full_inventory_in_appendix_a_is_accepted():
    from backend.consistency_mapping import locate_tables
    result = room_source()
    result['source_tables'][0]['chapter_path'] = ['附录 A 测评对象资产', '物理机房']
    result['source_tables'][0]['parent_titles'] = ['附录 A 测评对象资产', '物理机房']
    match = locate_tables(result, 'plan', 'full')['c01:plan:full']
    assert match['status'] == 'matched'


def test_unreadable_failed_table_is_not_reported_missing():
    from backend.consistency_mapping import locate_tables
    result = room_source()
    result['source_tables'][0]['parse_status'] = 'nested_table'
    result['source_tables'][0]['headers'] = []
    result['source_tables'][0]['cells'] = []
    match = locate_tables(result, 'plan', 'full')['c01:plan:full']
    assert match['status'] == 'parse_failed'
    assert match['table_ids'] == ['t1']


def test_policy_default_yields_to_actual_header_evidence():
    from backend.consistency_mapping import extract_cells

    def network_source(include_ip):
        doc = Document()
        doc.add_heading('2 测评对象', 1)
        doc.add_heading('2.4 系统构成', 2)
        doc.add_heading('网络设备', 3)
        headers = ['设备名称', '用途'] + (['IP地址'] if include_ip else [])
        table = doc.add_table(rows=2, cols=len(headers))
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ['设备 A', '接入'] + (['10.0.0.1'] if include_ip else [])):
            cell.text = value
        return parsed(doc)

    missing = [c for c in extract_cells(network_source(False), 'plan', 'full') if c['field_id'] == 'c02_f07']
    present = [c for c in extract_cells(network_source(True), 'plan', 'full') if c['field_id'] == 'c02_f07']
    assert len(missing) == len(present) == 1
    assert missing[0]['status'] == 'field_undefined'
    assert present[0]['status'] == 'value'
    assert present[0]['raw_value'] == '10.0.0.1'


def test_survey_checkbox_annotation_does_not_hide_virtual_device_column():
    from backend.consistency_mapping import extract_cells
    doc = Document()
    doc.add_heading('网络设备', 1)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['设备名称', '是否为虚拟设备（√/×）', '用途']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['设备 A', '√', '接入']):
        cell.text = value
    cells = [c for c in extract_cells(parsed(doc), 'survey', 'full') if c['field_id'] == 'c02_f02']
    assert len(cells) == 1
    assert cells[0]['status'] == 'value'
    assert cells[0]['raw_value'] == '√'


def comparison_cell(role, name, field='c01_f02', value='园区', status='value', row=2, category='c01'):
    return {'category_id': category, 'entity_name': name, 'field_id': field,
            'role': role, 'scope': 'full', 'raw_value': value, 'status': status,
            'source': {'table_id': 't1', 'row': row, 'col': 2, 'source_cell': f't1r{row}c2'}}


def compare_fixture(survey=None, plan=None, report=None, locations=None):
    return compare_documents({'survey': survey or [], 'plan': plan or [], 'report': report or []},
                             locations or {}, 'full')


def test_entity_key_preserves_case_and_parenthetical_text():
    result = compare_fixture([comparison_cell('survey', '设备(A)')],
                             [comparison_cell('plan', '设备 A')],
                             [comparison_cell('report', '设备a')])
    assert {row['entity_key'] for row in result['rows']} == {'设备A', '设备a'}
    assert len(result['rows']) == 2


def test_duplicate_key_retains_both_rows_without_cartesian_product():
    result = compare_fixture([comparison_cell('survey', '机房 A', value='园区甲', row=2),
                              comparison_cell('survey', '机房(A)', value='园区乙', row=3)],
                             [comparison_cell('plan', '机房A')],
                             [comparison_cell('report', '机房A')])
    assert len(result['rows']) == 1
    row = result['rows'][0]
    assert row['status'] == 'duplicate_key'
    assert [cell['source']['row'] for cell in row['source_records']['survey']] == [2, 3]
    assert row['fields']['c01_f02']['incomplete']
    issue = next(item for item in result['issues']
                 if item['code'] == 'duplicate_key' and item['field_id'] == 'c01_f02')
    assert [(cell['raw_value'], cell['source']['source_cell'])
            for cell in issue['evidence'] if cell['role'] == 'survey'] == [
                ('园区甲', 't1r2c2'), ('园区乙', 't1r3c2')]


def test_missing_object_empty_cell_unmapped_field_and_missing_table_differ():
    result = compare_fixture(
        [comparison_cell('survey', '甲', value='', status='empty'),
         comparison_cell('survey', '乙')],
        [comparison_cell('plan', '甲', value=None, status='field_undefined'),
         comparison_cell('plan', '乙')],
        [comparison_cell('report', '甲')],
        {'c01:report:full': {'status': 'matched'},
         'c01:plan:full': {'status': 'matched'},
         'c01:survey:full': {'status': 'matched'},
         'c02:report:full': {'status': 'missing', 'table_ids': []}})
    rows = {row['entity_key']: row for row in result['rows']}
    assert rows['甲']['fields']['c01_f02']['sources']['survey']['status'] == 'empty'
    assert rows['甲']['fields']['c01_f02']['sources']['plan']['status'] == 'field_undefined'
    assert rows['乙']['fields']['c01_f02']['sources']['report']['status'] == 'missing_object'
    assert any(d['status'] == 'missing' and d['position_id'] == 'c02:report:full' for d in result['diagnostics'])
    assert result['completeness'] is False


def test_address_compares_protocol_port_path():
    result = compare_fixture([comparison_cell('survey', '设备A', 'c02_f07', 'http://host:80/a', category='c02')],
                             [comparison_cell('plan', '设备A', 'c02_f07', 'http://host:80/a', category='c02')],
                             [comparison_cell('report', '设备A', 'c02_f07', 'https://host:443/b', category='c02')])
    field = result['rows'][0]['fields']['c02_f07']
    assert field['has_difference']
    assert field['sources']['report']['raw_value'] == 'https://host:443/b'


def test_crypto_model_and_certificate_stay_needs_review():
    result = compare_fixture([comparison_cell('survey', '产品A', 'c10_f02', '甲型', category='c10'),
                              comparison_cell('survey', '产品A', 'c10_f03', '证书甲', category='c10')],
                             [comparison_cell('plan', '产品A', 'c10_f02', '乙型', category='c10'),
                              comparison_cell('plan', '产品A', 'c10_f03', '证书乙', category='c10')],
                             [comparison_cell('report', '产品A', 'c10_f02', '甲型', category='c10'),
                              comparison_cell('report', '产品A', 'c10_f03', '证书甲', category='c10')])
    for field_id in ('c10_f02', 'c10_f03'):
        assert result['rows'][0]['fields'][field_id]['status'] == 'needs_review'


def test_report_baseline_missing_does_not_fall_back_to_plan():
    result = compare_fixture([comparison_cell('survey', '设备A', value='甲')],
                             [comparison_cell('plan', '设备A', value='乙')])
    field = result['rows'][0]['fields']['c01_f02']
    assert field['sources']['report']['status'] == 'missing_object'
    assert field['incomplete']
    assert field['has_difference'] is True
    assert set(field['different_sources']) == {'survey', 'plan', 'report'}


def test_missing_source_and_empty_cell_are_colored_as_report_differences():
    result = compare_fixture(
        [comparison_cell('survey', '设备A', value='', status='empty')],
        [],
        [comparison_cell('report', '设备A', value='园区')],
    )
    field = result['rows'][0]['fields']['c01_f02']
    assert field['has_difference'] is True
    assert set(field['different_sources']) == {'survey', 'plan'}
    assert field['sources']['plan']['status'] == 'missing_object'
    assert 'different_value' in field['issue_codes']


def test_matching_plan_stays_uncolored_when_only_survey_differs_from_report():
    result = compare_fixture(
        [comparison_cell('survey', '设备A', value='园区甲')],
        [comparison_cell('plan', '设备A', value='园区乙')],
        [comparison_cell('report', '设备A', value='园区乙')],
    )
    assert result['rows'][0]['fields']['c01_f02']['different_sources'] == ['survey']


def test_difference_and_incomplete_flags_can_coexist():
    result = compare_fixture([comparison_cell('survey', '设备A', value='甲')],
                             [comparison_cell('plan', '设备A', value='乙')],
                             [comparison_cell('report', '设备A', value=None, status='mapping_unresolved')])
    field = result['rows'][0]['fields']['c01_f02']
    assert field['has_difference'] is True
    assert field['incomplete'] is True
    assert result['completeness'] is False


def test_three_empty_values_are_equal_but_warn_empty():
    result = compare_fixture([comparison_cell('survey', '设备A', value='', status='empty')],
                             [comparison_cell('plan', '设备A', value='', status='empty')],
                             [comparison_cell('report', '设备A', value='', status='empty')])
    field = result['rows'][0]['fields']['c01_f02']
    assert field['has_difference'] is False
    assert field['status'] == 'empty'
    assert 'empty_value' in field['issue_codes']


def test_zero_and_false_are_values_not_missing():
    result = compare_fixture(
        [comparison_cell('survey', '设备A', value=0)],
        [comparison_cell('plan', '设备A', value=False)],
        [comparison_cell('report', '设备A', value=0)],
    )
    field = result['rows'][0]['fields']['c01_f02']
    assert field['has_difference']
    assert 'empty_value' not in field['issue_codes']


def test_checks_boundary_uses_located_cells_and_keeps_missing_tables_visible():
    from backend.checks import compare_consistency

    result = compare_consistency({
        'survey': {'status': 'parsed', 'parsed': {'source_tables': []}},
        'plan': {'status': 'parsed', 'parsed': {'source_tables': []}},
        'report': {'status': 'parsed', 'parsed': {'source_tables': []}},
    }, 'full')
    assert result['completeness'] is False
    assert any(d['status'] == 'missing' for d in result['diagnostics'])
