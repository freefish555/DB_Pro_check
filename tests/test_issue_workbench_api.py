import io
import os
import tempfile

os.environ.setdefault('REVIEW_DATA_DIR', tempfile.mkdtemp(prefix='issues-api-'))

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app, db_session, drain_tasks, issue_rule_version, startup
from backend.db import Audit, DATA, Document, Issue, Member, Project, Run, Session, Task, User


def setup_project():
    startup()
    with Session() as db:
        user = db.query(User).filter_by(username='admin').first()
        project = Project(name='=1+1', owner_id=user.id, config={})
        db.add(project)
        db.flush()
        db.add(Member(project_id=project.id, user_id=user.id, role='owner'))
        old = Run(project_id=project.id, request_key='old', mode='lenient', modules=[], snapshot={}, status='done', created_at=1)
        new = Run(project_id=project.id, request_key='new', mode='lenient', modules=[], snapshot={}, status='done', created_at=2)
        db.add_all([old, new])
        db.flush()
        db.add(Issue(project_id=project.id, run_id=old.id, chapter='CH01', category='writing', title='历史问题', description='old', suggestion='', evidence=[], machine={}, status='confirmed'))
        shared = Issue(project_id=project.id, run_id=new.id, chapter='CH01', category='writing', title='跨章问题', description='new', suggestion='', evidence=[], machine={'related_chapters': ['CH02'], 'rule_version': 'v2'}, status='confirmed')
        pending = Issue(project_id=project.id, run_id=new.id, chapter='CH02', category='writing', title='待确认问题', description='pending', suggestion='', evidence=[], machine={})
        db.add_all([shared, pending])
        db.commit()
        password = (DATA / 'bootstrap-admin.txt').read_text(encoding='utf-8').split('初始密码：')[1].splitlines()[0]
        return project.id, old.id, new.id, shared.id, pending.id, password


def login(client, password):
    csrf = client.post('/api/login', json={'username': 'admin', 'password': password}).json()['csrf']
    return {'X-CSRF-Token': csrf}


def test_current_run_and_related_chapters_share_one_issue_id():
    pid, old, new, shared, _, password = setup_project()
    with TestClient(app) as client:
        login(client, password)
        rows = client.get(f'/api/projects/{pid}/issues').json()
        assert {row['title'] for row in rows} == {'跨章问题', '待确认问题'}
        related = client.get(f'/api/projects/{pid}/issues?chapter=CH02').json()
        assert {row['id'] for row in related} == {shared, next(row['id'] for row in rows if row['title'] == '待确认问题')}
        assert client.get(f'/api/projects/{pid}/issues?run_id={old}').json()[0]['title'] == '历史问题'
        assert client.get(f'/api/projects/{pid}/issues?run_id=missing').status_code == 404


def test_report_content_is_available_before_any_review_run():
    _, _, _, _, _, password = setup_project()
    with Session() as db:
        user = db.query(User).filter_by(username='admin').first()
        project = Project(name='预处理验收', owner_id=user.id, config={})
        db.add(project); db.flush()
        db.add(Member(project_id=project.id, user_id=user.id, role='owner'))
        report = Document(project_id=project.id, role='report', filename='synthetic.docx',
                          sha256='synthetic', storage_key='synthetic', version=1, status='parsed',
                          parsed={'sections': [{'id': 's1', 'code': 'CH02'}],
                                  'blocks': [{'id': 'p1', 'type': 'paragraph', 'region': 'body',
                                              'section_id': 's1', 'text': '合成原文'}]})
        db.add(report); db.commit()
        pid, document_id = project.id, report.id
    with TestClient(app) as client:
        login(client, password)
        content = client.get(f'/api/projects/{pid}/issues/chapter-content?run_id=&chapter=CH02').json()
        assert content['document']['id'] == document_id
        assert content['blocks'][0]['text'] == '合成原文'
        assert content['review_status'] == '规则待配置'
        coverage = client.get(f'/api/projects/{pid}/issues/coverage?run_id=').json()
        assert next(row for row in coverage['chapters'] if row['code'] == 'CH02')['content_status'] == '已提取'


def test_batch_conflict_is_atomic_and_success_is_audited():
    pid, _, rid, shared, pending, password = setup_project()
    with TestClient(app) as client:
        headers = login(client, password)
        path = f'/api/projects/{pid}/issues/batch'
        body = {'run_id': rid, 'items': [{'id': shared, 'version': 1}, {'id': pending, 'version': 0}], 'status': 'rejected', 'note': '合成复核'}
        conflict = client.post(path, json=body, headers=headers)
        assert conflict.status_code == 409
        with Session() as db:
            assert db.get(Issue, shared).status == 'confirmed'
            assert db.get(Issue, pending).status == 'pending'
        body['items'][1]['version'] = 1
        response = client.post(path, json=body, headers=headers)
        assert response.status_code == 200, response.text
        assert {row['status'] for row in response.json()['items']} == {'rejected'}
        with Session() as db:
            assert {db.get(Issue, iid).version for iid in (shared, pending)} == {2}
            audits = db.query(Audit).filter_by(project_id=pid, action='issues_batch_reviewed').all()
            assert len(audits) == 1


def test_export_is_current_run_confirmed_only_and_states_zero_coverage():
    pid, old, new, shared, pending, password = setup_project()
    with TestClient(app) as client:
        login(client, password)
        result = client.get(f'/api/projects/{pid}/export?format=xlsx')
        assert result.status_code == 200
        book = load_workbook(io.BytesIO(result.content))
        assert book.active['C2'].value == '跨章问题'
        assert '历史问题' not in str(list(book.active.values))
        assert 'CH02' in str(list(book.active.values))
        assert book.active['J2'].value == 'v2'
        assert '审核覆盖情况' in book.sheetnames
        assert book['审核覆盖情况']['B1'].value == "'=1+1"
        doc = WordDocument(io.BytesIO(client.get(f'/api/projects/{pid}/export?format=docx').content))
        assert new in '\n'.join(p.text for p in doc.paragraphs)
        assert old not in '\n'.join(p.text for p in doc.paragraphs)
        with Session() as db:
            db.get(Issue, shared).status = 'rejected'
            db.commit()
        empty = load_workbook(io.BytesIO(client.get(f'/api/projects/{pid}/export?format=xlsx').content))
        assert empty.active.max_row == 1
        assert '0条已确认' in str(list(empty['审核覆盖情况'].values))
        coverage = client.get(f'/api/projects/{pid}/issues/coverage?run_id={new}').json()
        assert coverage['run_id'] == new
        assert coverage['confirmed_count'] == 0
        assert any(row['review_status'] == '规则待配置' for row in coverage['chapters'])


def test_batch_rejects_duplicate_ids_and_wrong_run():
    pid, old, new, shared, _, password = setup_project()
    with TestClient(app) as client:
        headers = login(client, password)
        path = f'/api/projects/{pid}/issues/batch'
        body = {'run_id': new, 'items': [{'id': shared, 'version': 1}, {'id': shared, 'version': 1}], 'status': 'confirmed'}
        assert client.post(path, json=body, headers=headers).status_code == 422
        body['items'] = [{'id': shared, 'version': 1}]
        body['run_id'] = old
        assert client.post(path, json=body, headers=headers).status_code == 409
        body['run_id'] = new
        body['items'] = [{'id': ['bad'], 'version': 1}]
        assert client.post(path, json=body, headers=headers).status_code == 422
        body['items'] = [{'id': shared, 'version': 1}]
        body['status'] = 'rejected'
        body['note'] = None
        assert client.post(path, json=body, headers=headers).status_code == 422
        with Session() as db:
            assert db.get(Issue, shared).version == 1


def test_single_review_audit_records_before_after_and_versions():
    pid, _, _, shared, _, password = setup_project()
    with TestClient(app) as client:
        headers = login(client, password)
        response = client.patch(f'/api/projects/{pid}/issues/{shared}',
                                json={'version': 1, 'status': 'needs_evidence', 'note': '待补凭据'},
                                headers=headers)
        assert response.status_code == 200
        with Session() as db:
            entry = db.query(Audit).filter_by(project_id=pid, action='issue_reviewed').one()
            assert entry.detail['before'] == {'status': 'confirmed', 'version': 1}
            assert entry.detail['status'] == 'needs_evidence'
            assert entry.detail['version'] == 2


def test_single_review_rejects_a_concurrent_version_change():
    pid, _, _, shared, _, password = setup_project()
    with TestClient(app) as client:
        headers = login(client, password)

        def racing_session():
            with Session() as db:
                original_execute = db.execute

                def execute(statement, *args, **kwargs):
                    if getattr(getattr(statement, 'table', None), 'name', None) == Issue.__table__.name:
                        with Session() as other:
                            other.query(Issue).filter_by(id=shared).update(
                                {'status': 'rejected', 'note': '另一位审核员', 'version': 2})
                            other.commit()
                    return original_execute(statement, *args, **kwargs)

                db.execute = execute
                yield db

        app.dependency_overrides[db_session] = racing_session
        try:
            response = client.patch(f'/api/projects/{pid}/issues/{shared}',
                                    json={'version': 1, 'status': 'confirmed'}, headers=headers)
        finally:
            app.dependency_overrides.pop(db_session, None)
        assert response.status_code == 409
        with Session() as db:
            row = db.get(Issue, shared)
            assert (row.status, row.note, row.version) == ('rejected', '另一位审核员', 2)
            assert not db.query(Audit).filter_by(project_id=pid, action='issue_reviewed').count()


def test_single_review_requires_reason_for_rejection_or_more_evidence():
    pid, _, _, shared, _, password = setup_project()
    with TestClient(app) as client:
        headers = login(client, password)
        for status in ('rejected', 'needs_evidence'):
            response = client.patch(f'/api/projects/{pid}/issues/{shared}',
                                    json={'version': 1, 'status': status, 'note': '   '}, headers=headers)
            assert response.status_code == 422
        with Session() as db:
            assert db.get(Issue, shared).version == 1


def test_coverage_uses_each_chapters_own_task_and_parse_status():
    pid, _, rid, _, _, password = setup_project()
    with Session() as db:
        run = db.get(Run, rid)
        report = Document(project_id=pid, role='report', filename='report.docx', sha256='sample',
                          storage_key='sample', version=1, status='partial', parsed={
                              'sections': [{'code': 'CH05'}, {'code': 'APP_D'}],
                              'source_tables': [{'location': 'APP_D · 表2', 'parse_status': 'nested_table'}],
                              'diagnostics': ['APP_D · 表2：嵌套表未解析']})
        db.add(report)
        db.flush()
        run.modules = ['high_risk', 'appendix_d']
        run.snapshot = {'document_ids': {'report': report.id}}
        run.status = 'partial'
        db.add_all([
            Task(project_id=pid, run_id=rid, kind='risk', label='risk', payload={}, status='done', output={}),
            Task(project_id=pid, run_id=rid, kind='record', label='record', payload={}, status='blocked', output={}),
        ])
        db.commit()
    with TestClient(app) as client:
        login(client, password)
        chapters = {row['code']: row for row in client.get(f'/api/projects/{pid}/issues/coverage?run_id={rid}').json()['chapters']}
        assert chapters['CH05']['review_status'] == '已完成'
        assert chapters['CH05']['content_status'] == '已提取'
        assert chapters['APP_D']['review_status'] == '部分失败'
        assert chapters['APP_D']['content_status'] == '提取不完整'


def test_all_status_export_labels_rows_and_records_filters():
    pid, _, rid, _, _, password = setup_project()
    with Session() as db:
        risk_task=Task(project_id=pid,run_id=rid,kind='risk',label='risk',payload={},status='done',output={})
        db.add(risk_task);db.flush()
        for issue in db.query(Issue).filter_by(run_id=rid):issue.task_id=risk_task.id
        db.commit()
    with TestClient(app) as client:
        login(client, password)
        url = f'/api/projects/{pid}/export?format=xlsx&run_id={rid}&module=high_risk&status=all'
        book = load_workbook(io.BytesIO(client.get(url).content))
        assert book.active.title == '全部问题'
        assert book.active['K1'].value == '复核状态'
        assert {row[10] for row in list(book.active.values)[1:]} == {'confirmed', 'pending'}
        coverage = dict((row[0], row[1]) for row in book['审核覆盖情况'].iter_rows(min_row=1, max_row=7, values_only=True))
        assert coverage['模块'] == 'high_risk'
        assert coverage['状态'] == 'all'
        doc = WordDocument(io.BytesIO(client.get(
            f'/api/projects/{pid}/export?format=docx&run_id={rid}&module=high_risk&status=all').content))
        assert '模块：high_risk' in doc.paragraphs[1].text
        assert '状态：all' in doc.paragraphs[1].text
        headings = [(p.text, p.style.name) for p in doc.paragraphs if p.style.name.startswith('Heading')]
        assert any(text.startswith('CH01') and style == 'Heading 1' for text, style in headings)
        assert any(text.startswith('CH02') and style == 'Heading 1' for text, style in headings)
        assert headings.index(next(h for h in headings if h[0].startswith('CH01'))) < headings.index(next(h for h in headings if h[0].startswith('CH02')))


def test_chapter_content_uses_run_document_snapshot_and_preserves_order():
    pid, _, rid, _, _, password = setup_project()
    with Session() as db:
        report = Document(project_id=pid, role='report', filename='reviewed.docx', sha256='old',
                          storage_key='old', version=1, status='parsed', parsed={
                              'sections': [{'id': 's1', 'code': 'CH02', 'title': '第2章', 'start_block': 'b1', 'end_block': 'b3'}],
                              'blocks': [
                                  {'id': 'b1', 'order': 1, 'type': 'paragraph', 'region': 'body', 'section_id': 's1',
                                   'text': '第2章', 'source': {'location': '正文 · 段落1'}},
                                  {'id': 'b2', 'order': 2, 'type': 'paragraph', 'region': 'body', 'section_id': 's1',
                                   'text': '原始段落', 'source': {'location': '正文 · 段落2'}},
                                  {'id': 'b3', 'order': 3, 'type': 'table', 'region': 'body', 'section_id': 's1',
                                   'table_id': 't1', 'text': '', 'source': {'location': '正文 · 表1'}},
                              ],
                              'source_tables': [{'table_id': 't1', 'block_id': 'b3', 'caption': '资产表', 'location': 'CH02 · 表1',
                                                 'parse_status': 'parsed', 'headers': ['名称'],
                                                 'cells': [{'row': 2, 'column': 1, 'raw_value': '测试设备'}]},
                                                {'table_id': 't1n1', 'parent_table_id': 't1', 'block_id': 'b3',
                                                 'caption': '嵌套信息', 'parse_status': 'parsed',
                                                 'cells': [{'row': 1, 'column': 1, 'raw_value': '内层单元格'}]}]})
        newer = Document(project_id=pid, role='report', filename='newer.docx', sha256='new',
                         storage_key='new', version=2, status='parsed', parsed={'sections': [], 'blocks': []})
        db.add_all([report, newer]); db.flush()
        run = db.get(Run, rid)
        run.snapshot = {'document_ids': {'report': report.id}}
        db.commit()
    with TestClient(app) as client:
        login(client, password)
        result = client.get(f'/api/projects/{pid}/issues/chapter-content?run_id={rid}&chapter=CH02')
        assert result.status_code == 200, result.text
        data = result.json()
        assert data['document']['version'] == 1
        assert [block['id'] for block in data['blocks']] == ['b1', 'b2', 'b3']
        assert data['blocks'][2]['table']['cells'][0]['raw_value'] == '测试设备'
        assert data['blocks'][2]['nested_tables_data'][0]['cells'][0]['raw_value'] == '内层单元格'
        assert data['review_status'] == '规则待配置'
        coverage = {item['code']: item for item in client.get(
            f'/api/projects/{pid}/issues/coverage?run_id={rid}').json()['chapters']}
        assert coverage['FULL_TEXT']['content_status'] == '已提取'
        assert coverage['APP_D']['content_status'] == '不适用/未出现'
        assert client.get(f'/api/projects/{pid}/issues/chapter-content?run_id={rid}&chapter=BOGUS').status_code == 422


def test_issue_rule_version_comes_from_the_immutable_run_snapshot():
    run = type('RunSnapshot', (), {'snapshot': {'mapping_version': 'map-v2',
                                                'algorithm_version': 'compare-v3',
                                                'guide_ids': ['guide-a'],
                                                'knowledge_ids': ['knowledge-b']}})()
    assert issue_rule_version(run, 'assets') == 'map-v2 / compare-v3'
    assert issue_rule_version(run, 'risk') == 'guide-a'
    assert issue_rule_version(run, 'record') == 'knowledge-b'
