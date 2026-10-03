"""Synthetic DOCX cases for the full-text source map."""

from base64 import b64decode
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.documents import CHAPTERS, parse_docx


def parsed(document, source_version=None):
    stream = BytesIO()
    document.save(stream)
    return parse_docx(stream.getvalue(), 'report', source_version=source_version)


def test_body_blocks_keep_word_order_and_toc_out_of_section_tree():
    document = Document()
    document.styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph('第三章 风险分析 …… 9', style='TOC 1')
    document.add_heading('第3章 风险分析', 1)
    first = document.add_paragraph('第一行')
    first.add_run().add_break()
    first.add_run('第二行')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '表格内容'
    document.add_paragraph('表格之后')
    document.add_heading('3.1 安全问题', 2)
    document.add_heading('等级测评结论', 1)
    result = parsed(document)

    assert [(b['type'], b.get('text', '')) for b in result['blocks']] == [
        ('toc', '第三章 风险分析 …… 9'),
        ('paragraph', '第3章 风险分析'),
        ('paragraph', '第一行\n第二行'),
        ('table', ''),
        ('paragraph', '表格之后'),
        ('paragraph', '3.1 安全问题'),
        ('paragraph', '等级测评结论'),
    ]
    assert result['blocks'][3]['table_id'] == result['source_tables'][0]['table_id']
    assert [s['title'] for s in result['sections']] == [
        '第3章 风险分析', '3.1 安全问题', '等级测评结论',
    ]
    assert result['sections'][1]['parent_id'] == result['sections'][0]['id']
    assert result['sections'][2]['parent_id'] is None
    assert result['sections'][0]['start_block'] == result['blocks'][1]['id']
    assert result['sections'][0]['end_block'] == result['blocks'][5]['id']
    assert result['sections'][2]['code'] == 'CONCLUSION'
    assert result['sections'][2]['number_status'] == 'not_applicable'
    assert result['blocks'][0]['section_id'] is None


def test_automatic_heading_number_is_preserved_as_evidence():
    document = Document()
    root = document.part.numbering_part.element
    abstract_id = max(int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))) + 1
    num_id = max(int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for tag, value in (('start', '1'), ('numFmt', 'decimal'), ('lvlText', '第%1章')):
        child = OxmlElement(f'w:{tag}')
        child.set(qn('w:val'), value)
        level.append(child)
    abstract.append(level)
    root.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), str(abstract_id))
    num.append(reference)
    root.append(num)
    heading = document.add_heading('安全通信网络', 1)
    num_properties = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    num_ref = OxmlElement('w:numId')
    num_ref.set(qn('w:val'), str(num_id))
    num_properties.extend((ilvl, num_ref))
    heading._p.get_or_add_pPr().append(num_properties)

    section = parsed(document)['sections'][0]
    assert section['number'] == '第1章'
    assert section['number_status'] == 'resolved'
    assert section['number_source'] == 'automatic'
    assert section['code'] == 'CH01'


def test_nested_table_and_picture_have_locators_and_diagnostics():
    document = Document()
    document.add_heading('第2章 测评对象', 1)
    outer = document.add_table(rows=1, cols=1)
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '嵌套内容'
    picture = document.add_paragraph()
    picture.add_run().add_picture(BytesIO(b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9WlZ5OQAAAAASUVORK5CYII='
    )))
    result = parsed(document)

    table_block = next(b for b in result['blocks'] if b['type'] == 'table')
    image_block = next(b for b in result['blocks'] if b['type'] == 'paragraph' and b.get('images'))
    assert table_block['table_id'] == 't1'
    assert table_block['nested_tables'][0]['location'].startswith('t1r1c1')
    assert result['source_tables'][1]['parent_table_id'] == 't1'
    assert result['source_tables'][1]['cells'][0]['raw_value'] == '嵌套内容'
    assert image_block['images'][0]['location'].startswith(image_block['source']['location'])
    assert image_block['images'][0]['relationship_id']
    assert image_block['section_id'] == result['sections'][0]['id']
    assert any('嵌套表' in item for item in result['diagnostics'])
    assert any('图片' in item for item in result['diagnostics'])


def test_textbox_is_located_without_being_mistaken_for_an_image():
    document = Document()
    document.add_heading('第2章 测评对象', 1)
    drawing = OxmlElement('w:drawing')
    box = OxmlElement('w:txbxContent')
    paragraph = OxmlElement('w:p')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '框内文字'
    run.append(text)
    paragraph.append(run)
    box.append(paragraph)
    drawing.append(box)
    document.add_paragraph()._p.append(drawing)

    result = parsed(document)
    block = result['blocks'][1]
    assert block['textboxes'][0]['text'] == '框内文字'
    assert block['images'] == []
    assert any('文本框' in item for item in result['diagnostics'])


def test_changed_reupload_has_distinct_document_source_references():
    document = Document()
    document.add_paragraph('旧报告')
    old = parsed(document)
    document.paragraphs[0].text = '新报告'
    new = parsed(document)

    old_ref = old['blocks'][0]['source']
    new_ref = new['blocks'][0]['source']
    assert old_ref['document_sha256'] != new_ref['document_sha256']
    assert old_ref['block_id'] == new_ref['block_id']
    assert old_ref['location'] == new_ref['location'] == '正文 · 段落1'


def test_identical_reupload_can_be_disambiguated_by_source_version():
    document = Document()
    document.add_paragraph('同一份报告')
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()

    first = parse_docx(data, 'report', source_version=1)['blocks'][0]['source']
    second = parse_docx(data, 'report', source_version=2)['blocks'][0]['source']
    assert first['document_sha256'] == second['document_sha256']
    assert first['document_version'] == 1
    assert second['document_version'] == 2
    assert first != second


def test_explicit_chapter_number_without_space_is_preserved():
    document = Document()
    document.add_heading('第4章整体测评', 1)
    section = parsed(document)['sections'][0]
    assert section['code'] == 'CH04'
    assert section['number'] == '第4章'
    assert section['number_source'] == 'explicit'


def test_unnumbered_report_regions_have_named_sections():
    document = Document()
    for title in ('等级保护测评报告', '基本信息表', '声明', '整改建议'):
        document.add_paragraph(title)
    result = parsed(document)
    assert [s['code'] for s in result['sections']] == [
        'COVER', 'BASIC_INFO', 'STATEMENT', 'RECTIFICATION',
    ]
    assert all(s['number_status'] == 'not_applicable' for s in result['sections'])
    assert all(s['code'] in dict(CHAPTERS) for s in result['sections'])


def test_header_and_footer_text_stay_out_of_body_section():
    document = Document()
    document.sections[0].header.paragraphs[0].text = '页眉 公司名'
    document.sections[0].footer.paragraphs[0].text = '页脚 第 1 页'
    document.add_heading('第1章 概述', 1)
    result = parsed(document)

    assert [b['region'] for b in result['blocks']] == ['body', 'header', 'footer']
    assert result['blocks'][1]['text'] == '页眉 公司名'
    assert result['blocks'][2]['text'] == '页脚 第 1 页'
    assert all(b['section_id'] is None for b in result['blocks'][1:])


def test_explicit_chapter_in_normal_style_opens_section():
    document = Document()
    document.add_paragraph('第4章 整体测评')
    document.add_paragraph('本章内容')

    result = parsed(document)
    assert [(section['code'], section['number']) for section in result['sections']] == [('CH04', '第4章')]
    assert result['blocks'][1]['section_id'] == result['sections'][0]['id']


def test_wrapped_body_paragraphs_and_tables_keep_document_order():
    document = Document()
    body = document.element.body
    first = document.add_paragraph('封面之前')
    wrapped = document.add_paragraph('第2章 测评对象')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '表内内容'
    last = document.add_paragraph('表格之后')
    for tag, element in [('w:sdt', wrapped._p), ('w:customXml', table._tbl)]:
        body.remove(element)
        wrapper = OxmlElement(tag)
        content = OxmlElement('w:sdtContent') if tag == 'w:sdt' else wrapper
        content.append(element)
        if content is not wrapper:
            wrapper.append(content)
        body.insert(1 if tag == 'w:sdt' else 2, wrapper)

    result = parsed(document)
    assert [(b['type'], b['text']) for b in result['blocks']] == [
        ('paragraph', first.text), ('paragraph', wrapped.text), ('table', ''), ('paragraph', last.text),
    ]
    assert result['source_tables'][0]['cells'][0]['raw_value'] == '表内内容'
    assert [section['code'] for section in result['sections']] == ['OTHER', 'CH02']
    assert result['blocks'][0]['section_id'] == result['sections'][0]['id']
    assert result['blocks'][2]['section_id'] == result['sections'][1]['id']


def test_altchunk_reports_unread_external_content():
    document = Document()
    document.add_paragraph('前文')
    chunk = OxmlElement('w:altChunk')
    chunk.set(qn('r:id'), 'rId999')
    document.element.body.insert(1, chunk)

    result = parsed(document)
    assert any('altChunk' in message and 'rId999' in message for message in result['diagnostics'])


def test_first_and_even_page_headers_and_footers_are_extracted():
    document = Document()
    document.add_paragraph('正文')
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    for label, container in [
        ('首页页眉', section.first_page_header), ('首页页脚', section.first_page_footer),
        ('偶数页眉', section.even_page_header), ('偶数页脚', section.even_page_footer),
    ]:
        container.paragraphs[0].text = label

    result = parsed(document)
    assert {(b['region'], b['text']) for b in result['blocks'][1:]} == {
        ('header', '首页页眉'), ('footer', '首页页脚'),
        ('header', '偶数页眉'), ('footer', '偶数页脚'),
    }
    assert all(b['section_id'] is None for b in result['blocks'][1:])


def test_image_in_table_has_row_and_cell_locator():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(1, 1).paragraphs[0].add_run().add_picture(BytesIO(b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9WlZ5OQAAAAASUVORK5CYII='
    )))

    result = parsed(document)
    image = result['blocks'][0]['images'][0]
    assert 't1r2c2' in image['location']
    assert image['relationship_id']


def test_unrecognized_top_level_heading_resets_previous_chapter():
    document = Document()
    document.add_heading('第2章 测评对象', 1)
    document.add_heading('未知一级标题', 1)
    document.add_paragraph('需人工归类')

    result = parsed(document)
    assert [section['code'] for section in result['sections']] == ['CH02', 'OTHER']
    assert result['blocks'][2]['section_id'] == result['sections'][1]['id']
    assert any('未知一级标题' in message and '未归类' in message for message in result['diagnostics'])
