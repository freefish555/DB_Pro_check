import io
import os
import tempfile
from pathlib import Path

os.environ['REVIEW_DATA_DIR']=tempfile.mkdtemp(prefix='db-pro-test-')

import pytest
import httpx
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app import app
from backend.ai import evaluate
from backend.db import Session, Issue, Knowledge, ModelConfig
from backend.catalog import import_workbook, select_requirements, selection_summary
from backend.checks import compare_assets, match_requirement, risk_review
from backend.documents import parse_docx


def word(role, name='核心交换机1', ip='10.0.0.1'):
    doc=WordDocument()
    if role=='report':
        doc.add_paragraph('附录 A 资产清单',style='Heading 1')
        doc.add_paragraph('附录A 表-1 网络设备')
    elif role=='plan':doc.add_paragraph('表 2-3 网络设备')
    else:doc.add_paragraph('表1-7 网络设备')
    table=doc.add_table(rows=1,cols=4)
    for i,x in enumerate(['序号','设备名称','IP地址','重要程度']):table.rows[0].cells[i].text=x
    cells=table.add_row().cells
    for i,x in enumerate(['1',name,ip,'关键']):cells[i].text=x
    if role=='plan':
        doc.add_paragraph('表 3-2 网络设备')
        second=doc.add_table(rows=1,cols=4)
        for i,x in enumerate(['序号','设备名称','IP地址','重要程度']):second.rows[0].cells[i].text=x
        cells=second.add_row().cells
        for i,x in enumerate(['1',name,ip,'关键']):cells[i].text=x
        doc.add_paragraph('表 4-1 测评工具')
        tool=doc.add_table(rows=1,cols=4)
        for i,x in enumerate(['序号','设备名称','型号','使用人']):tool.rows[0].cells[i].text=x
        cells=tool.add_row().cells
        for i,x in enumerate(['1','测试笔记本','T14','测试员']):cells[i].text=x
    if role=='report':
        doc.add_paragraph('单项测评结果记录',style='Heading 1')
        doc.add_paragraph('安全通信网络',style='Heading 2')
        doc.add_paragraph(name,style='Heading 4')
        record=doc.add_table(rows=1,cols=4)
        for i,x in enumerate(['控制点','测评项','结果记录','符合情况']):record.rows[0].cells[i].text=x
        cells=record.add_row().cells
        for i,x in enumerate(['网络架构','应保证关键网络设备业务处理能力','经检查，设备容量足够。','符合']):cells[i].text=x
        doc.add_paragraph('风险分析',style='Heading 1')
        risk=doc.add_table(rows=1,cols=4)
        for i,x in enumerate(['安全问题','关联资产','危害分析结果','风险等级']):risk.rows[0].cells[i].text=x
        cells=risk.add_row().cells
        for i,x in enumerate(['设备容量不足',name,'故判为中风险','低']):cells[i].text=x
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()


def catalog(level,rows):
    book=Workbook();sheet=book.active;sheet.title='安全通信网络'
    sheet.append(['序号','控制点','SAG属性标识','测评项','核查点数','核查点','判定规则'])
    for n,sag in enumerate(rows,1):sheet.append([n,'网络架构',sag,'应保证关键网络设备业务处理能力',1,'1）检查设备容量','满足为符合；否则为不符合'])
    stream=io.BytesIO();book.save(stream)
    return import_workbook(stream.getvalue(),f'安全通用要求_{"二级" if level==2 else "三级"}_核查点梳理统计表格.xlsx')


def test_realistic_document_scopes_and_evidence():
    survey=parse_docx(word('survey'),'survey');plan=parse_docx(word('plan'),'plan');report=parse_docx(word('report'),'report')
    assert len(plan['assets'])==2  # measurement tool is excluded
    assert {a['scope'] for a in plan['assets']}=={'full','planned_sample'}
    assert report['records'][0]['source'].startswith('附录D')
    result=compare_assets({r:{'status':'parsed','parsed':p} for r,p in [('survey',survey),('plan',plan),('report',report)]},'full')
    assert result['rows'][0]['status']=='consistent'
    assert len(result['issues'])==0
    mismatch=parse_docx(word('report',ip='10.0.0.2'),'report')
    result=compare_assets({r:{'status':'parsed','parsed':p} for r,p in [('survey',survey),('plan',plan),('report',mismatch)]},'full')
    assert any(i['category']=='asset_difference' for i in result['issues'])
    risk=risk_review(report,[{'scenario':'设备容量不足','key':'test','clause':'1'}])
    assert any(i['category']=='risk_consistency' for i in risk['issues'])


def test_sag_selection_and_manual_match():
    low=catalog(2,['S']);high=catalog(3,['A','G'])
    cats=[dict(id=str(i),family='general',level=level,profile='',content=x) for i,(level,x) in enumerate([(2,low),(3,high)])]
    selected=select_requirements(cats,{'s':2,'a':3,'g':3,'extensions':[]})
    assert selection_summary(selected)['by_attribute']=={'S':1,'A':1,'G':1}
    report=parse_docx(word('report'),'report')
    assert match_requirement(report['records'][0],selected)['status']=='ambiguous'  # duplicate text needs reviewer
    assert match_requirement(report['records'][0],selected,selected[0]['key'])['status']=='matched'


def test_auth_upload_run_and_review():
    with TestClient(app) as client:
        password=(Path(os.environ['REVIEW_DATA_DIR'])/'bootstrap-admin.txt').read_text(encoding='utf-8').split('初始密码：')[1].splitlines()[0]
        response=client.post('/api/login',json={'username':'admin','password':password})
        assert response.status_code==200
        csrf=response.json()['csrf'];headers={'X-CSRF-Token':csrf}
        assert client.post('/api/projects',json={'name':'测试项目'}).status_code==403
        project=client.post('/api/projects',json={'name':'测试项目'},headers=headers).json()
        for role in ('survey','plan','report'):
            response=client.post(f'/api/projects/{project["id"]}/documents',data={'role':role},files={'file':(role+'.docx',word(role))},headers=headers)
            assert response.status_code==200,response.text
        checked=client.post(f'/api/projects/{project["id"]}/precheck',json={'modules':['assets_full']},headers=headers).json()
        assert not checked['blockers']
        run=client.post(f'/api/projects/{project["id"]}/runs',json={'modules':['assets_full'],'mode':'lenient','request_key':'test-1'},headers=headers)
        assert run.status_code==200,run.text
        same=client.post(f'/api/projects/{project["id"]}/runs',json={'modules':['assets_full'],'request_key':'test-1'},headers=headers)
        assert same.json()['id']==run.json()['id']
        issue_rows=client.get(f'/api/projects/{project["id"]}/issues').json()
        assert issue_rows==[]
        with Session() as db:
            row=Issue(project_id=project['id'],chapter='APP_D',category='writing',title='合成问题',description='=1+1',suggestion='人工修正',object_name='核心交换机1',evidence=[{'source':'附录D表1行2','quote':'合成原文'}],machine={})
            db.add(row);db.commit();iid=row.id
        current=client.get(f'/api/projects/{project["id"]}/issues?chapter=APP_D').json()[0]
        updated=client.patch(f'/api/projects/{project["id"]}/issues/{iid}',json={'version':current['version'],'status':'confirmed','note':'已核对'},headers=headers)
        assert updated.status_code==200
        assert client.patch(f'/api/projects/{project["id"]}/issues/{iid}',json={'version':current['version'],'status':'rejected'},headers=headers).status_code==409
        exported=client.get(f'/api/projects/{project["id"]}/export?format=xlsx')
        assert exported.status_code==200
        from openpyxl import load_workbook
        sheet=load_workbook(io.BytesIO(exported.content)).active
        assert sheet['D2'].value=="'=1+1" and sheet['H2'].value=='已核对'


def test_model_contract_rejects_fabricated_quote(monkeypatch):
    import json
    original=httpx.Client
    point={'id':'p1','text':'核对容量'}
    record={'text':'已检查设备容量。','verdict':'符合'}
    requirement={'text':'应保证容量','points':[point],'decision':'满足为符合'}
    answer={'point_results':[{'point_id':'p1','coverage':'covered','quote':'不存在的原文','reason':''}],
            'verdict_review':{'status':'supported','reason':'','quote':''},'writing':[]}
    def handler(_):return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(answer,ensure_ascii=False)}}]})
    transport=httpx.MockTransport(handler)
    monkeypatch.setattr(httpx,'Client',lambda **kwargs: original(transport=transport,**kwargs))
    with pytest.raises(ValueError,match='引用'):
        evaluate(record,requirement,'strict',{'base_url':'http://127.0.0.1:9999/v1','model':'local','timeout':10},'')


def test_appendix_run_separates_three_issue_types(monkeypatch):
    import importlib
    import time
    module=importlib.import_module('backend.app')
    source=catalog(3,['G'])
    with Session() as db:
        k=Knowledge(filename='synthetic.xlsx',sha256='synthetic-appendix',family='general',level=3,profile='',status='published',content=source)
        db.add(k)
        db.add(ModelConfig(base_url='http://127.0.0.1:9999/v1',model='mock',enabled=True,external=False,timeout=10))
        db.commit()
    def fake(record,requirement,mode,model,key):
        assert mode=='strict' and requirement['decision']
        return {'point_results':[{'point_id':requirement['points'][0]['id'],'coverage':'missing','quote':'','reason':'缺少检查过程'}],
                'verdict_review':{'status':'contradicted','reason':'结论缺证据','quote':''},
                'writing':[{'type':'unclear','quote':'设备容量足够','suggestion':'写出检查依据'}]}
    monkeypatch.setattr(module,'evaluate',fake)
    with TestClient(app) as client:
        password=(Path(os.environ['REVIEW_DATA_DIR'])/'bootstrap-admin.txt').read_text(encoding='utf-8').split('初始密码：')[1].splitlines()[0]
        csrf=client.post('/api/login',json={'username':'admin','password':password}).json()['csrf']
        headers={'X-CSRF-Token':csrf}
        project=client.post('/api/projects',json={'name':'附录D合成项目'},headers=headers).json()
        upload=client.post(f'/api/projects/{project["id"]}/documents',data={'role':'report'},files={'file':('report.docx',word('report'))},headers=headers)
        assert upload.status_code==200
        run=client.post(f'/api/projects/{project["id"]}/runs',json={'modules':['appendix_d'],'mode':'strict','request_key':'ai-test'},headers=headers)
        assert run.status_code==200,run.text
        until=time.time()+5
        while time.time()<until:
            state=client.get(f'/api/projects/{project["id"]}/runs/{run.json()["id"]}').json()
            if state['tasks'][0]['status']=='done':break
            time.sleep(.05)
        assert state['tasks'][0]['status']=='done',state
        categories={x['category'] for x in client.get(f'/api/projects/{project["id"]}/issues?chapter=APP_D').json()}
        assert categories=={'description_coverage','verdict','writing'}

