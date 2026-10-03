"""Intranet review application. Run with: uvicorn backend.app:app"""
import hashlib
import hmac
import io
import ipaddress
import json
import os
import secrets
import time
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urlparse

from cryptography.fernet import Fernet
from docx import Document as WordDocument
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, select, update
from sqlalchemy.orm import Session as DbSession

from .ai import evaluate
from .appendix import classify_appendix
from .catalog import import_workbook, select_requirements, selection_summary
from .checks import LEGACY_RISK_COLUMNS, compare_assets, compare_consistency, issue, legacy_risk_rows, match_requirement, risk_review
from .consistency_mapping import MAPPING
from .db import (DATA, Base, engine, Session, User, LoginSession, Project, Member, Document,
                 Knowledge, ModelConfig, ModelService, UserModelProfile, Run, Task, RedactionBatch, RedactionApproval, ReviewDecision, Issue, Audit, public)
from .documents import CHAPTERS, parse_docx
from .key_info import build_key_info, _field_candidates
from .model_profiles import public_profile, profile_snapshot, profile_is_current
from .redaction import prepare_model_request, assert_authorized

app = FastAPI(title='等保测评报告评审平台', docs_url=None, redoc_url=None)
POOL = ThreadPoolExecutor(max_workers=4)
COOKIE = 'review_session'
ROLES = {'survey', 'plan', 'report'}
MODULES = {'assets_full', 'assets_sample', 'appendix_d', 'high_risk'}
CONSISTENCY_ALGORITHM_VERSION = 'consistency-v1'
CONSISTENCY_CATEGORIES = [{'id':c['id'],'label':c['label'],
                           'fields':[{'id':f['id'],'label':f['label']} for f in c['fields']]}
                          for c in MAPPING['categories']]


def appendix_dynamic_payload(record, requirement):
    """Build the complete dynamic object that crosses the redaction boundary."""
    return {
        'record': {**record, 'record_id': record.get('id', '')},
        'requirement': dict(requirement),
    }


def project_redaction_terms(configured, documents):
    """Include known project identities so unlabeled prose is redacted too."""
    terms={key:list(values) if isinstance(values,(list,tuple)) else [values]
           for key,values in (configured or {}).items()}
    kinds={'name':('tested_leader','tested_contact','assessor_contact','author','reviewer','approver'),
           'organization':('tested_name','assessor_name'),
           'address':('tested_address','assessor_address'),
           'phone':('tested_leader_phone','tested_phone','assessor_phone'),
           'identifier':('filing_number','credit_code','organization_code'),
           'asset':('system_name',)}
    fields={field_id:kind for kind,ids in kinds.items() for field_id in ids}
    for role,document in documents.items():
        candidates,_=_field_candidates({'id':document.id,'version':document.version,
                                         'sha256':document.sha256,'parsed':document.parsed or {}},role)
        for field_id,kind in fields.items():
            terms.setdefault(kind,[]).extend(item['value'] for item in candidates[field_id]
                                             if len(item['value'].strip())>=2)
        terms.setdefault('name',[]).extend(asset['name'] for asset in (document.parsed or {}).get('assets',[])
                                           if asset.get('type')=='person' and len(asset.get('name','').strip())>=2)
    return {kind:list(dict.fromkeys(value for value in values if isinstance(value,str) and value.strip()))
            for kind,values in terms.items()}


def password_hash(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    value = hashlib.pbkdf2_hmac('sha256', password.encode(), salt, 250_000)
    return salt.hex() + ':' + value.hex()


def password_valid(password, saved):
    salt, _ = saved.split(':', 1)
    return hmac.compare_digest(password_hash(password, bytes.fromhex(salt)), saved)


def audit(db, user, action, project=None, **detail):
    db.add(Audit(user_id=user.id if user else None, project_id=project, action=action, detail=detail))


def invalidate_redaction_authorizations(db, project_id):
    """Revoke previews and approvals when the project redaction dictionary changes."""
    runs = db.scalars(select(Run).where(Run.project_id == project_id)).all()
    for run in runs:
        if 'appendix_d' not in (run.modules or []):
            continue
        preview = (run.snapshot or {}).get('redaction_preview')
        if not preview:
            continue
        db.query(RedactionApproval).filter(RedactionApproval.run_id == run.id).delete(
            synchronize_session=False
        )
        for task in db.scalars(select(Task).where(Task.run_id == run.id, Task.kind == 'record')).all():
            # A running worker will fail its approval check before the model call.
            # Clear only non-running payloads so a later refresh rebuilds them.
            if task.status != 'running' and task.status not in ('done', 'cancelled'):
                task.status = 'awaiting_redaction'
                task.output = {}
        run.snapshot = {**(run.snapshot or {}), 'redaction_preview': {
            **preview, 'status': 'awaiting_confirmation', 'invalidated': True,
        }}
        run.status = 'awaiting_redaction'


def master_key():
    path = DATA / 'model.key'
    if not path.exists():
        path.write_bytes(Fernet.generate_key())
        try: os.chmod(path, 0o600)
        except OSError: pass
    return path.read_bytes()


@app.on_event('startup')
def startup():
    Base.metadata.create_all(engine)
    with Session() as db:
        if not db.scalar(select(User.id).limit(1)):
            secret = secrets.token_urlsafe(20)
            db.add(User(username='admin', display_name='管理员', password_hash=password_hash(secret), admin=True))
            db.commit()
            path = DATA / 'bootstrap-admin.txt'
            path.write_text('用户名：admin\n初始密码：' + secret + '\n首次登录后请修改。\n', encoding='utf-8')
            try: os.chmod(path, 0o600)
            except OSError: pass
    master_key()
    POOL.submit(recover_tasks)


def db_session():
    with Session() as db:
        yield db


def current(request: Request, db: DbSession = Depends(db_session)):
    token = request.cookies.get(COOKIE, '')
    if not token: raise HTTPException(401, '请先登录')
    sess = db.scalar(select(LoginSession).where(LoginSession.token_hash == hashlib.sha256(token.encode()).hexdigest()))
    if not sess or sess.expires_at < time.time(): raise HTTPException(401, '登录已过期')
    user = db.get(User, sess.user_id)
    if not user or not user.active: raise HTTPException(401, '用户不可用')
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        if request.headers.get('x-csrf-token') != sess.csrf: raise HTTPException(403, '请求校验失败')
        origin = request.headers.get('origin')
        if origin and urlparse(origin).netloc != request.headers.get('host'):
            raise HTTPException(403, '请求来源不匹配')
    return user, sess


def admin(auth=Depends(current)):
    if not auth[0].admin: raise HTTPException(403, '需要管理员权限')
    return auth[0]


def project_access(db, project_id, user):
    project = db.get(Project, project_id)
    if not project or (not user.admin and not db.scalar(select(Member.id).where(Member.project_id == project_id, Member.user_id == user.id))):
        raise HTTPException(404, '项目不存在或无权限')
    return project


def payload(request_data, keys):
    if not isinstance(request_data, dict) or any(k not in request_data for k in keys):
        raise HTTPException(422, '缺少必要字段')
    return request_data


@app.post('/api/login')
def login(body: dict, response: Response, request: Request, db=Depends(db_session)):
    origin = request.headers.get('origin')
    if origin and urlparse(origin).netloc != request.headers.get('host'): raise HTTPException(403, '请求来源不匹配')
    name = str(body.get('username', ''))[:100]
    user = db.scalar(select(User).where(User.username == name))
    if not user or not user.active or user.locked_until > time.time() or not password_valid(str(body.get('password', '')), user.password_hash):
        if user and user.active:
            user.failures += 1
            if user.failures >= 5: user.locked_until = time.time() + 900; user.failures = 0
            db.commit()
        raise HTTPException(401, '用户名或密码错误，连续失败会暂时锁定账号')
    user.failures = 0; user.locked_until = 0
    token = secrets.token_urlsafe(40); csrf = secrets.token_urlsafe(32)
    db.add(LoginSession(user_id=user.id, token_hash=hashlib.sha256(token.encode()).hexdigest(), csrf=csrf, expires_at=time.time()+8*3600))
    audit(db, user, 'login'); db.commit()
    response.set_cookie(COOKIE, token, httponly=True, secure=request.url.scheme == 'https', samesite='strict', max_age=8*3600)
    return {'user':public(user, ('password_hash','failures','locked_until')), 'csrf':csrf}


@app.get('/api/me')
def me(auth=Depends(current)):
    user, session = auth
    return {'user':public(user, ('password_hash','failures','locked_until')), 'csrf':session.csrf}


@app.post('/api/logout')
def logout(response: Response, auth=Depends(current), db=Depends(db_session)):
    db.delete(auth[1]); audit(db, auth[0], 'logout'); db.commit(); response.delete_cookie(COOKIE)
    return {'ok':True}


@app.post('/api/me/password')
def change_password(body: dict, auth=Depends(current), db=Depends(db_session)):
    user = db.get(User, auth[0].id)
    if not password_valid(str(body.get('old', '')), user.password_hash): raise HTTPException(422, '原密码错误')
    new = str(body.get('new', ''))
    if len(new) < 12: raise HTTPException(422, '新密码至少12位')
    user.password_hash = password_hash(new); audit(db,user,'password_changed'); db.commit()
    if user.username=='admin':(DATA / 'bootstrap-admin.txt').unlink(missing_ok=True)
    return {'ok':True}


@app.get('/api/users')
def users(_=Depends(admin), db=Depends(db_session)):
    return [public(x, ('password_hash','failures','locked_until')) for x in db.scalars(select(User).order_by(User.created_at))]


@app.post('/api/users')
def add_user(body: dict, user=Depends(admin), db=Depends(db_session)):
    payload(body, ('username','display_name','password'))
    if len(body['password']) < 12: raise HTTPException(422, '密码至少12位')
    if db.scalar(select(User.id).where(User.username==body['username'])): raise HTTPException(409,'用户名已存在')
    row=User(username=body['username'][:100], display_name=body['display_name'][:150],password_hash=password_hash(body['password']),admin=bool(body.get('admin',False)))
    db.add(row);audit(db,user,'user_created',user_id=row.username);db.commit()
    return public(row,('password_hash','failures','locked_until'))


@app.get('/api/projects')
def projects(auth=Depends(current), db=Depends(db_session)):
    user=auth[0]
    rows=db.scalars(select(Project).order_by(Project.created_at.desc())).all()
    return [public(p) for p in rows if user.admin or db.scalar(select(Member.id).where(Member.project_id==p.id,Member.user_id==user.id))]


@app.post('/api/projects')
def add_project(body: dict, auth=Depends(current), db=Depends(db_session)):
    name=str(body.get('name','')).strip()
    if not name or len(name)>200:raise HTTPException(422,'项目名称不能为空且不超过200字')
    p=Project(name=name, owner_id=auth[0].id,config={'s':3,'a':3,'g':3,'extensions':[],'power_category':None,'aliases':{},'matches':{}})
    db.add(p);db.flush();db.add(Member(project_id=p.id,user_id=auth[0].id,role='owner'))
    audit(db,auth[0],'project_created',p.id);db.commit();return public(p)


@app.get('/api/projects/{pid}')
def get_project(pid:str,auth=Depends(current),db=Depends(db_session)):
    return public(project_access(db,pid,auth[0]))


@app.patch('/api/projects/{pid}')
def edit_project(pid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0]); v=body.get('version')
    if v!=p.version:raise HTTPException(409,'项目已经被其他审核员修改，请刷新')
    config=dict(p.config or {})
    if 'name' in body:p.name=str(body['name']).strip()[:200]
    for k in ('s','a','g'):
        if k in body:
            value=int(body[k]);
            if value not in (2,3,4):raise HTTPException(422,'当前知识库仅支持2至4级')
            config[k]=value
    if config.get('g')!=max(config.get('s',3),config.get('a',3)):
        raise HTTPException(422,'依据附录A，G等级应为S与A等级中较高者')
    mapping=next((k for k in knowledge_rows(db) if k['family']=='mapping'),None)
    if mapping and {'s':config['s'],'a':config['a'],'g':config['g']} not in mapping['content'].get('profiles',[]):
        raise HTTPException(422,'该S/A/G组合不在已发布的附录A等级映射中')
    if 'extensions' in body:
        ext=body['extensions']
        if not isinstance(ext,list) or any(x not in ('cloud','mobile','iot','ics','bigdata') for x in ext):raise HTTPException(422,'扩展类型无效；电力类别请单独二选一')
        config['extensions']=list(dict.fromkeys(ext))
    if 'power_category' in body:
        power=body['power_category'] or None
        if power not in (None,'power_monitoring','power_management'):raise HTTPException(422,'电力类别必须二选一')
        config['power_category']=power
    for k in ('aliases','matches'):
        if k in body:
            if not isinstance(body[k],dict):raise HTTPException(422,f'{k} 应为对象')
            config[k]=body[k]
    if 'redaction_terms' in body:
        terms=body['redaction_terms']
        if not isinstance(terms,dict):raise HTTPException(422,'脱敏词典必须是对象')
        normalized={}
        total=0
        for key,values in terms.items():
            if not isinstance(key,str) or len(key)>60 or not isinstance(values,list):
                raise HTTPException(422,'脱敏词典格式无效')
            clean=[]
            for value in values:
                value=str(value).strip()
                if not value or len(value)>300:continue
                if value not in clean:clean.append(value)
            total+=len(clean)
            normalized[key]=clean[:1000]
        if total>10000:raise HTTPException(422,'脱敏词典条目过多')
        config['redaction_terms']=normalized
    redaction_changed = config.get('redaction_terms') != (p.config or {}).get('redaction_terms', {})
    p.config=config;p.version+=1
    if redaction_changed:
        invalidate_redaction_authorizations(db, pid)
    audit(db,auth[0],'project_updated',pid);db.commit();return public(p)


@app.post('/api/projects/{pid}/redaction-terms')
async def upload_redaction_terms(pid: str, file: UploadFile = File(...), auth=Depends(current), db: Session=Depends(db_session)):
    project_access(db, pid, auth[0])
    name, data = await office_bytes(file, '.xlsx')
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(422, '脱敏字段工作簿无法读取') from exc
    aliases = {'姓名': 'name', '人名': 'name', '单位': 'organization', '公司': 'organization',
               '地址': 'address', '电话': 'phone', '手机': 'phone', 'ip': 'ip', '域名': 'domain',
               '网址': 'url', 'url': 'url', '编号': 'identifier', '证书': 'identifier'}
    terms = {}
    for sheet in workbook.worksheets:
        for values in sheet.iter_rows(values_only=True):
            cells = [str(value).strip() for value in values if value is not None and str(value).strip()]
            if len(cells) < 2:
                continue
            kind = aliases.get(cells[0].casefold())
            if not kind:
                continue
            terms.setdefault(kind, []).extend(cells[1:])
    normalized = {key: list(dict.fromkeys(value))[:1000] for key, value in terms.items()}
    if not normalized:
        raise HTTPException(422, '工作簿中未找到可识别的脱敏字段')
    project = db.get(Project, pid); config = dict(project.config or {}); config['redaction_terms'] = normalized
    invalidate_redaction_authorizations(db, pid)
    project.config = config; project.version += 1
    audit(db, auth[0], 'redaction_terms_uploaded', pid, filename=name, categories=list(normalized))
    db.commit()
    return public(project)


@app.post('/api/projects/{pid}/members')
def add_member(pid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0])
    if not auth[0].admin and p.owner_id!=auth[0].id:raise HTTPException(403,'只有项目负责人可添加成员')
    user=db.scalar(select(User).where(User.username==body.get('username')))
    if not user:raise HTTPException(404,'用户不存在')
    if not db.scalar(select(Member.id).where(Member.project_id==pid,Member.user_id==user.id)):
        db.add(Member(project_id=pid,user_id=user.id,role='reviewer'))
    audit(db,auth[0],'member_added',pid,username=user.username);db.commit();return {'ok':True}


async def office_bytes(file:UploadFile, suffix:str):
    name=Path(file.filename or '').name
    if not name.lower().endswith(suffix):raise HTTPException(422,f'只接收{suffix}文件')
    data=await file.read(100*1024*1024+1)
    if len(data)>100*1024*1024:raise HTTPException(413,'文件超过100MB')
    return name,data


def latest_documents(db,pid):
    rows=db.scalars(select(Document).where(Document.project_id==pid).order_by(Document.version.desc())).all()
    return {x.role:x for x in rows[::-1]}


@app.get('/api/projects/{pid}/documents')
def documents(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    return [public(x,('parsed','storage_key')) | {'counts':x.parsed.get('counts',{}),'diagnostics':x.parsed.get('diagnostics',[])} for x in db.scalars(select(Document).where(Document.project_id==pid).order_by(Document.created_at.desc()))]


@app.post('/api/projects/{pid}/documents')
async def upload_document(pid:str,role:str=Form(...),file:UploadFile=File(...),auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    if role not in ROLES:raise HTTPException(422,'文档类型无效')
    name,data=await office_bytes(file,'.docx')
    version=max([x.version for x in db.scalars(select(Document).where(Document.project_id==pid,Document.role==role))],default=0)+1
    try: parsed=parse_docx(data,role,source_version=version)
    except Exception as exc:raise HTTPException(422,f'文档解析失败：{str(exc)[:200]}') from exc
    key=uuid.uuid4().hex+'.docx'; folder=DATA/'uploads';folder.mkdir(exist_ok=True)
    (folder/key).write_bytes(data)
    row=Document(project_id=pid,role=role,filename=name,sha256=hashlib.sha256(data).hexdigest(),storage_key=key,version=version,status='partial' if parsed['diagnostics'] else 'parsed',parsed=parsed)
    db.add(row);audit(db,auth[0],'document_uploaded',pid,role=role,version=version,sha256=row.sha256);db.commit()
    return public(row,('parsed','storage_key')) | {'counts':parsed['counts'],'diagnostics':parsed['diagnostics']}


@app.get('/api/projects/{pid}/facts')
def facts(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]); docs=latest_documents(db,pid)
    return {role:{'status':d.status,'facts':d.parsed.get('facts',{}),'counts':d.parsed.get('counts',{}),'diagnostics':d.parsed.get('diagnostics',[]),'sections':d.parsed.get('sections',[])} for role,d in docs.items()}


@app.get('/api/projects/{pid}/key-info')
def key_info(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    docs=latest_documents(db,pid)
    return build_key_info({role:{'id':d.id,'version':d.version,'sha256':d.sha256,
                                 'status':d.status,'parsed':d.parsed} for role,d in docs.items()})


@app.get('/api/projects/{pid}/assets')
def assets(pid:str,scope:str='full',auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0]); docs=latest_documents(db,pid)
    if scope not in ('full','sample'):raise HTTPException(422,'范围无效')
    try:return compare_assets({k:{'id':v.id,'status':v.status,'parsed':v.parsed} for k,v in docs.items()},scope,(p.config or {}).get('aliases'))
    except ValueError as exc:return {'scope':scope,'rows':[],'issues':[],'blocked':str(exc)}


def knowledge_rows(db):
    rows=db.scalars(select(Knowledge).where(Knowledge.status=='published').order_by(Knowledge.created_at.desc())).all()
    selected={}
    for k in rows:
        key=(k.family,k.level,k.profile)
        if key not in selected:selected[key]=k
    return [public(k,('content',)) | {'content':k.content} for k in selected.values()]


@app.get('/api/knowledge')
def knowledge(auth=Depends(current),db=Depends(db_session)):
    return [public(k,('content',)) | {'count':len(k.content.get('requirements',[])),'warnings':k.content.get('warnings',[])} for k in db.scalars(select(Knowledge).order_by(Knowledge.created_at.desc()))]


@app.post('/api/knowledge')
async def upload_knowledge(file:UploadFile=File(...),user=Depends(admin),db=Depends(db_session)):
    name,data=await office_bytes(file,'.xlsx'); sha=hashlib.sha256(data).hexdigest()
    existing=db.scalar(select(Knowledge).where(Knowledge.sha256==sha))
    if existing:return public(existing,('content',)) | {'count':len(existing.content.get('requirements',[])),'warnings':existing.content.get('warnings',[])}
    try:content=import_workbook(data,name)
    except Exception as exc:raise HTTPException(422,f'核查点文件解析失败：{str(exc)[:200]}') from exc
    k=Knowledge(filename=name,sha256=sha,family=content['family'],level=content['level'],profile=content['profile'],content=content,status='draft')
    db.add(k);audit(db,user,'knowledge_uploaded',filename=name,sha256=sha);db.commit()
    return public(k,('content',)) | {'count':len(content.get('requirements',[])),'warnings':content.get('warnings',[])}


@app.post('/api/knowledge/{kid}/publish')
def publish_knowledge(kid:str,user=Depends(admin),db=Depends(db_session)):
    k=db.get(Knowledge,kid)
    if not k:raise HTTPException(404,'依据不存在')
    if k.family!='mapping' and not k.content.get('requirements'):raise HTTPException(422,'依据无有效条目')
    if k.family in ('general','power_monitoring','power_management') and k.level:
        catalogs=[row for row in knowledge_rows(db) if (row['family'],row['level'],row['profile']) != (k.family,k.level,k.profile)]
        catalogs.append(public(k,('content',)) | {'content':k.content})
        powers=('power_monitoring','power_management') if k.family=='general' else (k.family,)
        for power in powers:
            try:select_requirements(catalogs,{'s':k.level,'a':k.level,'g':k.level,'extensions':[],'power_category':power})
            except ValueError as exc:raise HTTPException(422,f'知识来源冲突，不能发布：{exc}') from exc
    k.status='published';k.published_by=user.id;audit(db,user,'knowledge_published',family=k.family,sha256=k.sha256);db.commit()
    return {'ok':True}


@app.get('/api/projects/{pid}/requirements')
def requirements(pid:str,auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0])
    try:reqs=select_requirements(knowledge_rows(db),p.config or {})
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    return {'summary':selection_summary(reqs),'items':reqs}


@app.get('/api/projects/{pid}/records')
def records(pid:str,auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0]); docs=latest_documents(db,pid); report=docs.get('report')
    if not report:return {'items':[],'blocked':'请先上传测评报告'}
    try:reqs=select_requirements(knowledge_rows(db),p.config or {})
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    overrides=(p.config or {}).get('matches',{})
    return {'items':[{'record':r,'match_key':f'{report.id}:{r["id"]}','match':match_requirement(r,reqs,overrides.get(f'{report.id}:{r["id"]}'))} for r in report.parsed.get('records',[])], 'count':len(report.parsed.get('records',[]))}


@app.get('/api/model')
def get_model(user=Depends(admin),db=Depends(db_session)):
    m=db.scalar(select(ModelConfig).order_by(ModelConfig.created_at.desc()))
    return {'configured':bool(m),'base_url':m.base_url if m else '', 'model':m.model if m else '', 'external':m.external if m else False,'enabled':m.enabled if m else False,'timeout':m.timeout if m else 120,'has_key':bool(m.key_encrypted) if m else False}


def model_service_input(body):
    url=str(body.get('base_url','')).strip().rstrip('/')
    parsed=urlparse(url)
    if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise HTTPException(422,'模型服务地址必须为无凭据的 HTTP(S) 地址')
    models=body.get('models')
    if not isinstance(models,list) or not models or any(not isinstance(x,str) or not x.strip() or len(x)>150 for x in models):
        raise HTTPException(422,'需配置可选模型标识')
    name=str(body.get('name','')).strip()[:150]
    if not name:raise HTTPException(422,'模型服务名称不能为空')
    return name,url,list(dict.fromkeys(x.strip() for x in models))


@app.get('/api/model-services')
def model_services(auth=Depends(current),db=Depends(db_session)):
    return [public(row) for row in db.scalars(select(ModelService).order_by(ModelService.created_at))]


@app.post('/api/model-services')
def add_model_service(body:dict,user=Depends(admin),db=Depends(db_session)):
    name,url,models=model_service_input(body)
    if db.scalar(select(ModelService.id).where(ModelService.base_url==url)):
        raise HTTPException(409,'该服务地址已经获准')
    row=ModelService(name=name,base_url=url,models=models,external=bool(body.get('external')),
                     enabled=bool(body.get('enabled',True)))
    db.add(row);audit(db,user,'model_service_created',service_id=row.id);db.commit()
    return public(row)


@app.patch('/api/model-services/{sid}')
def change_model_service(sid:str,body:dict,user=Depends(admin),db=Depends(db_session)):
    row=db.get(ModelService,sid)
    if not row:raise HTTPException(404,'模型服务不存在')
    if 'base_url' in body or 'models' in body or 'name' in body:
        name,url,models=model_service_input({'name':body.get('name',row.name),
                                             'base_url':body.get('base_url',row.base_url),
                                             'models':body.get('models',row.models)})
        row.name=name;row.base_url=url;row.models=models
    if 'enabled' in body:row.enabled=bool(body['enabled'])
    if 'external' in body:row.external=bool(body['external'])
    row.version+=1;audit(db,user,'model_service_updated',service_id=sid);db.commit()
    return public(row)


def require_model_choice(db, service_id, model):
    service=db.get(ModelService,str(service_id))
    if not service or not service.enabled or model not in service.models:
        raise HTTPException(422,'模型服务未获准、已停用或模型标识不在允许列表')
    return service


@app.get('/api/me/model-profiles')
def my_model_profiles(auth=Depends(current),db=Depends(db_session)):
    return [public_profile(row) for row in db.scalars(select(UserModelProfile)
            .where(UserModelProfile.owner_id==auth[0].id).order_by(UserModelProfile.created_at))]


@app.post('/api/me/model-profiles')
def add_model_profile(body:dict,auth=Depends(current),db=Depends(db_session)):
    if set(body)-{'service_id','model','api_key','enabled'}:
        raise HTTPException(422,'个人档案只能选择管理员批准的服务')
    model=str(body.get('model','')).strip()
    service=require_model_choice(db,body.get('service_id'),model)
    key=str(body.get('api_key',''))
    if not key:raise HTTPException(422,'API Key 不能为空')
    row=UserModelProfile(owner_id=auth[0].id,service_id=service.id,model=model,
                         key_encrypted=Fernet(master_key()).encrypt(key.encode()).decode(),
                         enabled=bool(body.get('enabled',True)))
    db.add(row);audit(db,auth[0],'personal_model_created',profile_id=row.id);db.commit()
    return public_profile(row)


@app.patch('/api/me/model-profiles/{profile_id}')
def change_model_profile(profile_id:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    row=db.get(UserModelProfile,profile_id)
    if not row or row.owner_id!=auth[0].id:raise HTTPException(404,'模型档案不存在')
    if set(body)-{'service_id','model','api_key','enabled'}:
        raise HTTPException(422,'个人档案只能选择管理员批准的服务')
    service_id=body.get('service_id',row.service_id)
    model=str(body.get('model',row.model)).strip()
    require_model_choice(db,service_id,model)
    row.service_id=service_id;row.model=model
    if 'api_key' in body:
        key=str(body['api_key'])
        if not key:raise HTTPException(422,'API Key 不能为空')
        row.key_encrypted=Fernet(master_key()).encrypt(key.encode()).decode()
    if 'enabled' in body:row.enabled=bool(body['enabled'])
    row.version+=1;audit(db,auth[0],'personal_model_updated',profile_id=row.id);db.commit()
    return public_profile(row)


@app.put('/api/model')
def put_model(body:dict,user=Depends(admin),db=Depends(db_session)):
    url=str(body.get('base_url','')).strip(); parsed=urlparse(url)
    if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise HTTPException(422,'模型地址必须为完整 HTTP(S) 地址，不得包含凭据或查询参数')
    try: private_host=ipaddress.ip_address(parsed.hostname).is_private
    except ValueError: private_host=bool(parsed.hostname and ('.' not in parsed.hostname or parsed.hostname.endswith(('.local','.internal'))))
    if not body.get('external') and parsed.hostname!='localhost' and not private_host:
        raise HTTPException(422,'外部模型请明确勾选“允许发送文档内容到外部服务”')
    name=str(body.get('model','')).strip()
    if not name:raise HTTPException(422,'模型名称不能为空')
    m=db.scalar(select(ModelConfig).order_by(ModelConfig.created_at.desc()))
    if not m:m=ModelConfig(base_url=url,model=name);db.add(m)
    m.base_url=url;m.model=name;m.external=bool(body.get('external'));m.enabled=bool(body.get('enabled',True));m.timeout=min(max(int(body.get('timeout',120)),10),300)
    if body.get('api_key'):m.key_encrypted=Fernet(master_key()).encrypt(str(body['api_key']).encode()).decode()
    audit(db,user,'model_config_updated',external=m.external,host=parsed.hostname);db.commit()
    return get_model(user,db)


def precheck(db,pid,modules,user=None,profile_id=None,allow_redaction_preview=False):
    docs=latest_documents(db,pid); p=db.get(Project,pid); blockers=[]; counts={}
    for role in ('survey','plan','report'):
        counts[role]=docs[role].parsed.get('counts',{}) if role in docs else None
    if any(m in modules for m in ('assets_full','assets_sample')):
        for module,scope in [('assets_full','full'),('assets_sample','sample')]:
            if module not in modules:continue
            for role in (('survey','plan','report') if scope=='full' else ('plan','report')):
                doc=docs.get(role)
                if not doc or doc.status not in ('parsed','partial') or 'source_tables' not in (doc.parsed or {}):
                    blockers.append(f'{module}：{role} 文档缺失或尚无可用的来源表解析快照')
    try:reqs=select_requirements(knowledge_rows(db),p.config or {})
    except ValueError as exc:
        blockers.append(f'核查点选择冲突：{exc}');reqs=[]
    if 'appendix_d' in modules:
        if not allow_redaction_preview:
            blockers.append('附录D：请先创建脱敏预览，再确认请求级授权')
        report=docs.get('report')
        if not report or not report.parsed.get('records'):blockers.append('附录D：尚未识别结果记录')
        if not reqs:blockers.append('附录D：当前等级和扩展未选出核查点')
        profile=db.get(UserModelProfile,profile_id) if profile_id else None
        service=db.get(ModelService,profile.service_id) if profile else None
        if not user or not profile or profile.owner_id!=user.id or not profile.enabled or not service or not service.enabled or profile.model not in service.models:
            blockers.append('附录D：请选择本人已启用且管理员批准的模型档案')
    if 'high_risk' in modules:
        report=docs.get('report')
        if not report or not report.parsed.get('risk_tables'):blockers.append('高风险：未识别报告中的问题、整体测评或风险表')
        if not any(k['family']=='high_risk' for k in knowledge_rows(db)):blockers.append('高风险：尚未发布判定指引')
    return {'blockers':blockers,'counts':counts,'requirements':selection_summary(reqs)}


@app.post('/api/projects/{pid}/precheck')
def check_run(pid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]); modules=body.get('modules',[])
    if not modules or any(m not in MODULES for m in modules):raise HTTPException(422,'审核模块无效')
    return precheck(db,pid,modules,auth[0],body.get('model_profile_id'),bool(body.get('preview_only')))


@app.post('/api/projects/{pid}/runs')
def start_run(pid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);modules=body.get('modules',[])
    if not modules or any(m not in MODULES for m in modules):raise HTTPException(422,'审核模块无效')
    mode=body.get('mode','lenient')
    if mode not in ('lenient','strict'):raise HTTPException(422,'审核模式无效')
    request_key=str(body.get('request_key') or uuid.uuid4())[:100]
    existing=db.scalar(select(Run).where(Run.project_id==pid,Run.request_key==request_key))
    if existing:return public(existing)
    preview_only=bool(body.get('preview_only'))
    check=precheck(db,pid,modules,auth[0],body.get('model_profile_id'),preview_only)
    if check['blockers']:raise HTTPException(422,{'blockers':check['blockers']})
    p=db.get(Project,pid);docs=latest_documents(db,pid);ks=knowledge_rows(db)
    reqs=select_requirements(ks,p.config or {})
    profile=db.get(UserModelProfile,body.get('model_profile_id')) if 'appendix_d' in modules else None
    service=db.get(ModelService,profile.service_id) if profile else None
    snapshot={'project_config':p.config,'document_ids':{r:d.id for r,d in docs.items()},
              'document_versions':{r:{'id':d.id,'version':d.version,'sha256':d.sha256} for r,d in docs.items()},
              'mapping_version':MAPPING['mapping_version'],'algorithm_version':CONSISTENCY_ALGORITHM_VERSION,
              'consistency_categories':CONSISTENCY_CATEGORIES,
              'created_at':time.time(),'knowledge_ids':list({kid for r in reqs for kid in r.get('knowledge_ids',[r['knowledge_id']])}),
              'selected_requirements':reqs,
              'guide_ids':[k['id'] for k in ks if k['family']=='high_risk'],
              'model_profile':profile_snapshot(profile,service) if profile else None,
              'created_by':auth[0].id}
    run=Run(project_id=pid,request_key=request_key,mode=mode,modules=modules,snapshot=snapshot,status='queued')
    db.add(run);db.flush()
    task_rows=[]
    def task(kind,label,payload):
        row=Task(project_id=pid,run_id=run.id,kind=kind,label=label,payload=payload,status='queued')
        db.add(row);db.flush();task_rows.append(row);return row
    if 'assets_full' in modules:task('assets','三文档完整资产核查',{'scope':'full'})
    if 'assets_sample' in modules:task('assets','方案与报告抽选对象核查',{'scope':'sample'})
    if 'high_risk' in modules:task('risk','高风险表格交叉核查',{})
    if 'appendix_d' in modules:
        report=docs['report'];overrides=(p.config or {}).get('matches',{})
        terms=project_redaction_terms((p.config or {}).get('redaction_terms',{}),docs)
        token_map={}
        for rec in report.parsed['records']:
            match=match_requirement(rec,reqs,overrides.get(f'{report.id}:{rec["id"]}'))
            row=task('record',f'{rec["object"]} · {rec["source"]}',{'record_id':rec['id'],'requirement_key':match['requirement']['key'] if match['status']=='matched' else None,'match_status':match['status']})
            if preview_only and match['status']=='matched':
                req=match['requirement']; context={'prompt_version':'appendix-v1','knowledge_version':','.join(sorted({str(x.get('knowledge_id','')) for x in reqs})),'service_url':service.base_url,'model_config_version':f'{profile.id}:{profile.version}:{service.version}'}
                prepared=prepare_model_request(appendix_dynamic_payload(rec,req),terms,token_map,context)
                # Keep one batch map so the same source value receives the same
                # token across every record preview in this run.
                token_map.update(prepared.token_map)
                row.status='awaiting_redaction';row.output={'redaction':{'payload':prepared.payload,'hash':prepared.request_hash,'uncertain':prepared.uncertain,'hits':prepared.hits}}
            elif preview_only:
                row.status='blocked'
                row.output={'record':rec,'reason':'结果记录尚未人工匹配核查点，补充匹配后可重新生成脱敏预览'}
    if preview_only and 'appendix_d' in modules:
        hashes=[]
        for row in task_rows:
            if row.kind!='record' or not row.output.get('redaction'):continue
            hashes.append({'task_id':row.id,'request_hash':row.output['redaction']['hash']})
        hashes.sort(key=lambda item:item['task_id'])
        list_hash=hashlib.sha256(json.dumps(hashes,sort_keys=True).encode()).hexdigest()
        snapshot['redaction_preview']={'hashes':hashes,'list_hash':list_hash,'status':'awaiting_confirmation'}
        run.status='awaiting_redaction'
        db.add(RedactionBatch(run_id=run.id,list_hash=list_hash,
                              token_map_encrypted=Fernet(master_key()).encrypt(json.dumps(token_map,ensure_ascii=False).encode()).decode(),
                              terms_encrypted=Fernet(master_key()).encrypt(json.dumps(terms,ensure_ascii=False).encode()).decode()))
    audit(db,auth[0],'run_started',pid,run_id=run.id,modules=modules,mode=mode);db.commit()
    # Preview mode only pauses Appendix D record tasks; consistency and
    # high-risk tasks may still run in the same multi-module run.
    for _ in range(4):POOL.submit(drain_tasks)
    return public(run)


@app.get('/api/projects/{pid}/runs/{rid}/redaction')
def redaction_preview(pid:str,rid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    rows=[]
    for task in db.scalars(select(Task).where(Task.run_id==rid,Task.kind=='record')):
        if task.status != 'awaiting_redaction':
            continue
        item=(task.output or {}).get('redaction')
        if item:rows.append({'task_id':task.id,'request_hash':item['hash'],'payload':item['payload'],'uncertain':item['uncertain'],'hits':item['hits'],'status':task.status})
    preview=run.snapshot.get('redaction_preview',{})
    return {'run_id':rid,'status':run.status,'list_hash':preview.get('list_hash',''),
            'items':rows,'counts':{'requests':len(rows),'uncertain':sum(bool(x['uncertain']) for x in rows),
                                    'hits':sum(len(x['hits']) for x in rows)}}


def _refresh_redaction_preview(db, pid, run, reviewer):
    """Rebuild pending Appendix D requests after an explicit version change."""
    project=project_access(db,pid,reviewer)
    snapshot=run.snapshot or {}
    profile_info=snapshot.get('model_profile') or {}
    profile=db.get(UserModelProfile,profile_info.get('profile_id'))
    service=db.get(ModelService,profile.service_id) if profile else None
    if not profile or profile.owner_id!=reviewer.id or not profile.enabled or not service or not service.enabled or profile.model not in service.models:
        raise HTTPException(422,'当前模型档案不可用，请重新选择本人已启用的模型档案')
    docs={role:db.get(Document,did) for role,did in snapshot.get('document_ids',{}).items()}
    report=docs.get('report')
    if not report or not report.parsed.get('records'):
        raise HTTPException(422,'运行快照中没有可用的附录D记录')
    terms=project_redaction_terms((project.config or {}).get('redaction_terms',{}),docs)
    token_map={}
    hashes=[]
    records={row['id']:row for row in report.parsed.get('records',[])}
    for task in db.scalars(select(Task).where(Task.run_id==run.id,Task.kind=='record')).all():
        if task.status in ('done','running'):
            continue
        reqkey=(task.payload or {}).get('requirement_key')
        if not reqkey:
            override=(project.config or {}).get('matches',{}).get(f"{report.id}:{(task.payload or {}).get('record_id')}")
            if override:
                reqkey=override
                task.payload={**(task.payload or {}),'requirement_key':override,'match_status':'manual_at_refresh'}
        record=records.get((task.payload or {}).get('record_id'))
        req=next((item for item in snapshot.get('selected_requirements',[]) if item.get('key')==reqkey),None)
        if not record or not req:
            task.status='blocked'
            task.output={'record':record} if record else {}
            continue
        context={'prompt_version':'appendix-v1',
                 'knowledge_version':','.join(sorted({str(x.get('knowledge_id','')) for x in snapshot.get('selected_requirements',[])})),
                 'service_url':service.base_url,
                 'model_config_version':f'{profile.id}:{profile.version}:{service.version}'}
        prepared=prepare_model_request(appendix_dynamic_payload(record,req),terms,token_map,context)
        token_map.update(prepared.token_map)
        task.status='awaiting_redaction';task.error=''
        task.output={'redaction':{'payload':prepared.payload,'hash':prepared.request_hash,
                                  'uncertain':prepared.uncertain,'hits':prepared.hits}}
        hashes.append({'task_id':task.id,'request_hash':prepared.request_hash})
    hashes.sort(key=lambda item:item['task_id'])
    list_hash=hashlib.sha256(json.dumps(hashes,sort_keys=True).encode()).hexdigest()
    snapshot={**snapshot,'model_profile':profile_snapshot(profile,service),
              'redaction_preview':{'hashes':hashes,'list_hash':list_hash,'status':'awaiting_confirmation'}}
    run.snapshot=snapshot;run.status='awaiting_redaction'
    batch=db.scalar(select(RedactionBatch).where(RedactionBatch.run_id==run.id))
    encrypted_map=Fernet(master_key()).encrypt(json.dumps(token_map,ensure_ascii=False).encode()).decode()
    encrypted_terms=Fernet(master_key()).encrypt(json.dumps(terms,ensure_ascii=False).encode()).decode()
    if batch:
        batch.version+=1;batch.list_hash=list_hash;batch.token_map_encrypted=encrypted_map;batch.terms_encrypted=encrypted_terms
        batch.reviewed_by=None;batch.reviewed_at=None
    else:
        db.add(RedactionBatch(run_id=run.id,version=1,list_hash=list_hash,
                              token_map_encrypted=encrypted_map,terms_encrypted=encrypted_terms))
    audit(db,reviewer,'redaction_preview_refreshed',pid,run_id=run.id,list_hash=list_hash)
    db.commit()
    return {'run_id':run.id,'status':run.status,'list_hash':list_hash,
            'items':[{'task_id':task.id,'request_hash':(task.output or {}).get('redaction',{}).get('hash'),
                      'payload':(task.output or {}).get('redaction',{}).get('payload'),
                      'uncertain':(task.output or {}).get('redaction',{}).get('uncertain',[]),
                      'hits':(task.output or {}).get('redaction',{}).get('hits',[]),
                      'status':task.status}
                     for task in db.scalars(select(Task).where(Task.run_id==run.id,Task.kind=='record')).all()
                     if task.status == 'awaiting_redaction' and (task.output or {}).get('redaction')]}


@app.post('/api/projects/{pid}/runs/{rid}/redaction/refresh')
def refresh_redaction(pid:str,rid:str,auth=Depends(current),db:Session=Depends(db_session)):
    run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    if 'appendix_d' not in (run.modules or []):raise HTTPException(422,'该运行不包含附录D')
    # A preview with unmatched records may settle as ``partial`` while other
    # modules finish; manual matching must still be able to refresh it.
    if run.status not in ('awaiting_redaction','queued','partial'):
        raise HTTPException(409,'redaction preview cannot be refreshed in the current run state')
    return _refresh_redaction_preview(db,pid,run,auth[0])


@app.post('/api/projects/{pid}/runs/{rid}/redaction/authorize')
def authorize_redaction(pid:str,rid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    preview=run.snapshot.get('redaction_preview',{})
    # Other modules in a mixed run may temporarily move the run to ``running``;
    # the immutable preview state is the authorization gate for Appendix D.
    if preview.get('status') != 'awaiting_confirmation':
        raise HTTPException(409,'����Ԥ���ѱ�Ȩ��������ˢ�µ�ǰԤ��')
    if preview.get('list_hash') and body.get('list_hash') not in (None,preview['list_hash']):
        raise HTTPException(409,'redaction preview changed')
    supplied=body.get('hashes'); supplied=supplied if isinstance(supplied,dict) else {x.get('task_id'):x.get('request_hash') for x in supplied or [] if isinstance(x,dict)}
    tasks=db.scalars(select(Task).where(Task.run_id==rid,Task.kind=='record')).all()
    pending_tasks=[task for task in tasks if task.status == 'awaiting_redaction' and (task.output or {}).get('redaction')]
    if not pending_tasks:
        raise HTTPException(422,'当前没有可授权的脱敏请求，请先人工匹配记录并刷新预览')
    approvals=[]
    for task in pending_tasks:
        item=(task.output or {}).get('redaction')
        if not item:continue
        if item['uncertain']:raise HTTPException(422,'仍有未确认的疑似敏感片段，不能放行')
        if supplied.get(task.id)!=item['hash']:raise HTTPException(409,'脱敏请求已变化或尚未确认，请重新预览')
        approval=db.scalar(select(RedactionApproval).where(RedactionApproval.run_id==rid,RedactionApproval.task_id==task.id))
        if approval:
            approval.request_hash=item['hash'];approval.reviewer_id=auth[0].id;approval.list_hash=preview.get('list_hash','')
        else:db.add(RedactionApproval(run_id=rid,task_id=task.id,request_hash=item['hash'],reviewer_id=auth[0].id,list_hash=preview.get('list_hash','')))
        task.status='queued';approvals.append(task.id)
    batch=db.scalar(select(RedactionBatch).where(RedactionBatch.run_id==rid))
    if batch:
        batch.reviewed_by=auth[0].id;batch.reviewed_at=time.time();batch.list_hash=preview.get('list_hash','')
    run.status='queued';run.snapshot={**run.snapshot,'redaction_preview':{**preview,'status':'authorized'}};db.commit()
    for _ in range(4):POOL.submit(drain_tasks)
    return {'ok':True,'authorized':approvals}


@app.get('/api/projects/{pid}/runs')
def runs(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);rows=db.scalars(select(Run).where(Run.project_id==pid).order_by(Run.created_at.desc())).all()
    return [{**public(r),'tasks':dict(Counter(t.status for t in db.scalars(select(Task).where(Task.run_id==r.id))))} for r in rows]


@app.get('/api/projects/{pid}/runs/{rid}')
def run_detail(pid:str,rid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);r=db.get(Run,rid)
    if not r or r.project_id!=pid:raise HTTPException(404,'任务不存在')
    return {'run':public(r),'tasks':[public(t) for t in db.scalars(select(Task).where(Task.run_id==rid).order_by(Task.created_at))]}


APPENDIX_COLUMNS=['序号','章节号','层面','测评对象','控制点','测评项','结果记录','符合情况','核查点','符合性判定与结果记录不一致','结果记录与核查点不相关','语句问题','错别字','核查点覆盖不足']


def appendix_snapshot(db,pid,rid,user,view='issues'):
    project_access(db,pid,user);run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    docs={role:db.get(Document,did) for role,did in run.snapshot.get('document_ids',{}).items()};report=docs.get('report')
    records={x['id']:x for x in (report.parsed or {}).get('records',[])} if report else {}
    tasks=db.scalars(select(Task).where(Task.run_id==rid,Task.kind=='record')).all();issue_rows=db.scalars(select(Issue).where(Issue.run_id==rid,Issue.chapter=='APP_D')).all()
    grouped={}
    for row in issue_rows:grouped.setdefault((row.machine or {}).get('record_id'),[]).append(row)
    rows=[]
    for task in tasks:
        record=(task.output or {}).get('record') or records.get((task.payload or {}).get('record_id'))
        if not record:continue
        found=grouped.get(record['id'],[])
        if view=='issues' and (not found or all(x.status == 'rejected' for x in found)):continue
        if view=='pending' and not any(x.status in ('pending','needs_evidence') for x in found):continue
        if view=='rejected' and not any(x.status=='rejected' for x in found):continue
        if view=='confirmed' and not any(x.status=='confirmed' for x in found):continue
        req=(task.output or {}).get('requirement') or {};issues={}
        for item in found:
            if view=='confirmed' and item.status!='confirmed':continue
            category = 'description_coverage' if item.category == 'review_pending' else item.category
            issues.setdefault(category,[]).append({'id':item.id,'text':item.description,'status':item.status,'version':item.version,'evidence':item.evidence,'note':item.note,'category':item.category})
        rows.append({'record_id':record['id'],'chapter_number':record.get('chapter_number','章节号待定位'),'layer':record.get('domain',''),'layer_order':record.get('layer_order',10**6),'object':record.get('object','对象待定位'),'object_order':record.get('object_order',10**6),'source_order':record.get('source_order',10**6),'control':record.get('control',''),'requirement':record.get('requirement',''),'text':record.get('text',''),'verdict':record.get('verdict',''),'points':req.get('points',[]),'issues':issues})
    rows.sort(key=lambda x:(x['layer_order'],x['object_order'],x['source_order'],x['record_id']))
    for number,row in enumerate(rows,1):row['number']=number
    return {'run_id':rid,'view':view,'columns':APPENDIX_COLUMNS,'rows':rows,'counts':{'records':len(rows),'issues':sum(len(v) for row in rows for v in row['issues'].values())}}


@app.get('/api/projects/{pid}/runs/{rid}/appendix')
def appendix_view(pid:str,rid:str,view:str='issues',auth=Depends(current),db=Depends(db_session)):
    if view not in ('issues','all','pending','rejected','confirmed'):raise HTTPException(422,'附录D视图无效')
    return appendix_snapshot(db,pid,rid,auth[0],view)


@app.get('/api/projects/{pid}/runs/{rid}/appendix/export')
def appendix_export(pid:str,rid:str,view:str='confirmed',auth=Depends(current),db=Depends(db_session)):
    if view not in ('confirmed','candidates'):raise HTTPException(422,'附录D导出视图无效')
    result=appendix_snapshot(db,pid,rid,auth[0],'confirmed' if view=='confirmed' else 'issues');book=Workbook();sheet=book.active;sheet.title='附录D结果记录核查';detail=book.create_sheet('问题复核明细');append_text_row(sheet,result['columns'])
    for row in result['rows']:
        def text(category):return '；'.join(x['text'] for x in row['issues'].get(category,[]))
        append_text_row(sheet,[row['number'],row['chapter_number'],row['layer'],row['object'],row['control'],row['requirement'],row['text'],row['verdict'],'；'.join(x.get('text','') for x in row['points']),text('verdict'),text('unrelated'),text('grammar'),text('typo'),text('description_coverage')])
        for category,items in row['issues'].items():
            for item in items:append_text_row(detail,[row['record_id'],category,item['status'],item['text'],json.dumps(item['evidence'],ensure_ascii=False),item['note']])
    detail.insert_rows(1);append_text_row(detail,['结果记录ID','问题列','复核状态','问题说明','原文定位','审核意见']);stream=io.BytesIO();book.save(stream);stream.seek(0);audit(db,auth[0],'appendix_exported',pid,run_id=rid,view=view,count=len(result['rows']));db.commit()
    return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="appendix-{view}.xlsx"'})


def consistency_result(db, pid, rid, user, scope, category='', status=''):
    project_access(db,pid,user)
    run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    if scope not in ('full','sample'):raise HTTPException(422,'范围无效')
    task=db.scalar(select(Task).where(Task.run_id==rid,Task.kind=='assets',
                                      Task.payload['scope'].as_string()==scope))
    if not task:raise HTTPException(404,'本次运行未选择该范围')
    if task.status!='done':raise HTTPException(409,'一致性核查尚未完成')
    output=task.output or {}
    if not output.get('mapping_version') or any('id' not in row for row in output.get('rows',[])):
        raise HTTPException(409,'旧运行没有新版一致性快照，请新建运行')
    decisions={row.row_id:public(row) for row in db.scalars(select(ReviewDecision).where(ReviewDecision.run_id==rid))}
    rows=[{**row,'review':decisions.get(row['id'])} for row in output.get('rows',[])]
    if category:rows=[row for row in rows if row['category_id']==category]
    if status:
        rows=[row for row in rows if row['status']==status or
              (status=='different' and row.get('has_difference')) or
              (status=='incomplete' and row.get('incomplete')) or
              (status=='parse_failed' and any(cell.get('status')=='parse_failed'
                  for field in row.get('fields',{}).values()
                  for cell in field.get('sources',{}).values())) or
              (row['review'] or {}).get('status')==status]
    ids={row['id'] for row in rows}
    issues=[item for item in output.get('issues',[]) if item.get('row_id') in ids]
    diagnostics=[item for item in output.get('diagnostics',[])
                 if not category or item.get('category_id')==category or item.get('position_id','').startswith(category+':')]
    return {'run_id':rid,'scope':scope,'mapping_version':output.get('mapping_version'),
            'algorithm_version':output.get('algorithm_version'),
            'categories':output.get('categories',[]),
            'document_versions':output.get('document_versions',run.snapshot.get('document_versions',{})),
            'created_at':output.get('created_at',run.created_at),
            'completeness':bool(output.get('completeness')) and not diagnostics,
            'rows':rows,'issues':issues,'diagnostics':diagnostics,
            'filters':{'scope':scope,'category':category,'status':status}}


@app.get('/api/projects/{pid}/runs/{rid}/consistency')
def get_consistency(pid:str,rid:str,scope:str='full',category:str='',status:str='',auth=Depends(current),db=Depends(db_session)):
    return consistency_result(db,pid,rid,auth[0],scope,category,status)


@app.patch('/api/projects/{pid}/runs/{rid}/consistency/{row_id}')
def review_consistency(pid:str,rid:str,row_id:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    run=db.get(Run,rid)
    if not run or run.project_id!=pid:raise HTTPException(404,'运行不存在')
    tasks=db.scalars(select(Task).where(Task.run_id==rid,Task.kind=='assets',Task.status=='done')).all()
    if not any(row.get('id')==row_id for task in tasks for row in (task.output or {}).get('rows',[])):
        raise HTTPException(404,'结果行不存在')
    status=body.get('status')
    if status not in ('pending','confirmed','rejected','needs_evidence'):
        raise HTTPException(422,'复核状态无效')
    existing=db.scalar(select(ReviewDecision).where(ReviewDecision.run_id==rid,ReviewDecision.row_id==row_id))
    expected=existing.version if existing else 0
    if body.get('version')!=expected:raise HTTPException(409,'结果已被其他审核员修改，请刷新')
    if existing:
        row=existing;row.version+=1
    else:
        row=ReviewDecision(run_id=rid,row_id=row_id,reviewer_id=auth[0].id,version=1)
        db.add(row)
    row.status=status;row.note=str(body.get('note',''))[:4000]
    row.reviewer_id=auth[0].id;row.updated_at=time.time()
    audit(db,auth[0],'consistency_reviewed',pid,run_id=rid,row_id=row_id,status=status)
    db.commit()
    return public(row)


def excel_text(value):
    value='' if value is None else str(value)
    return "'"+value if value.startswith(('=','+','-','@')) else value


def append_text_row(sheet, values):
    sheet.append([excel_text(value) for value in values])
    for cell in sheet[sheet.max_row]:cell.number_format='@'


@app.get('/api/projects/{pid}/runs/{rid}/consistency/export')
def export_consistency(pid:str,rid:str,scope:str='full',category:str='',status:str='',auth=Depends(current),db=Depends(db_session)):
    result=consistency_result(db,pid,rid,auth[0],scope,category,status)
    book=Workbook();wide=book.active;wide.title='三文档横向汇总'
    sheet=book.create_sheet('逐对象逐字段结果');exceptions=book.create_sheet('异常明细')
    append_text_row(wide,[f'筛选：范围={scope}；类别={category or "全部"}；状态={status or "全部"}'])
    roles=('survey','plan','report') if scope=='full' else ('plan','report')
    role_names={'survey':'调研表','plan':'测评方案','report':'测评报告'}
    categories=[item for item in MAPPING['categories'] if not category or item['id']==category]
    for item in categories:
        rows=[row for row in result['rows'] if row['category_id']==item['id']]
        fields=[(field['id'],field['label']) for field in item['fields']]
        known={field_id for field_id,_ in fields}
        for row in rows:
            for field in row.get('fields',{}).values():
                if field['field_id'] not in known:
                    fields.append((field['field_id'],field['label']))
                    known.add(field['field_id'])
        append_text_row(wide,[item['label']])
        append_text_row(wide,['是否一致','序号']+[value for _,label in fields for value in ([label]+['']*(len(roles)-1))]
                        +['对象键','行ID','人工状态','人工意见'])
        append_text_row(wide,['','']+[role_names[role] for _ in fields for role in roles]+['','','',''])
        for index,row in enumerate(rows,1):
            values=[row['status'],index]
            differing=[]
            for field_id,_ in fields:
                field=row.get('fields',{}).get(field_id,{})
                for role in roles:
                    values.append(field.get('sources',{}).get(role,{}).get('raw_value'))
                    if role in field.get('different_sources',[]):differing.append(len(values))
            review=row.get('review') or {}
            append_text_row(wide,values+[row['entity_key'],row['id'],review.get('status',''),review.get('note','')])
            for column in differing:wide.cell(wide.max_row,column).fill=PatternFill('solid',fgColor='00FFF2CC')
        wide.append([])
    wide.freeze_panes='C5'
    append_text_row(sheet,[f'筛选：范围={scope}；类别={category or "全部"}；状态={status or "全部"}'])
    append_text_row(sheet,['行ID','范围','类别','对象键','行结论','字段ID','字段','字段结论','来源','原值','源状态','出处','人工状态','人工意见'])
    for row in result['rows']:
        for field in row.get('fields',{}).values():
            for role,cell in field.get('sources',{}).items():
                source=cell.get('source')
                append_text_row(sheet,[row['id'],row['scope'],row['category'],row['entity_key'],row['status'],
                                       field['field_id'],field['label'],field['status'],role,cell.get('raw_value'),
                                       cell['status'],json.dumps(source,ensure_ascii=False) if source else '',
                                       (row['review'] or {}).get('status',''),(row['review'] or {}).get('note','')])
                if role in field.get('different_sources',[]):
                    sheet.cell(sheet.max_row,10).fill=PatternFill('solid',fgColor='00FFF2CC')
                    sheet.cell(sheet.max_row,11).fill=PatternFill('solid',fgColor='00FFF2CC')
    append_text_row(exceptions,['行ID','范围','类别','对象键','字段ID','异常代码','证据'])
    for item in result['issues']:
        append_text_row(exceptions,[item.get('row_id'),item.get('scope'),item.get('category_id'),
                                    item.get('entity_key'),item.get('field_id'),item.get('code'),
                                    json.dumps(item.get('evidence',[]),ensure_ascii=False)])
    for item in result['diagnostics']:
        append_text_row(exceptions,['',scope,item.get('category_id',''),'','',item.get('status'),
                                    json.dumps(item,ensure_ascii=False)])
    stream=io.BytesIO();book.save(stream);stream.seek(0)
    audit(db,auth[0],'consistency_exported',pid,run_id=rid,scope=scope,category=category,status=status)
    db.commit()
    return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition':'attachment; filename="consistency.xlsx"'})


@app.post('/api/projects/{pid}/runs/{rid}/cancel')
def cancel_run(pid:str,rid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);r=db.get(Run,rid)
    if not r or r.project_id!=pid:raise HTTPException(404,'任务不存在')
    r.cancelled=True;r.status='cancelled'
    for t in db.scalars(select(Task).where(Task.run_id==rid,Task.status=='queued')):t.status='cancelled'
    audit(db,auth[0],'run_cancelled',pid,run_id=rid);db.commit();return {'ok':True}


@app.post('/api/projects/{pid}/tasks/{tid}/retry')
def retry_task(pid:str,tid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);t=db.get(Task,tid)
    if not t or t.project_id!=pid:raise HTTPException(404,'子任务不存在')
    if t.status not in ('failed','blocked'):raise HTTPException(409,'仅失败或待确认任务可重试')
    r=db.get(Run,t.run_id)
    if t.kind=='record' and not t.payload.get('requirement_key'):
        report_id=r.snapshot['document_ids']['report']
        override=(db.get(Project,pid).config or {}).get('matches',{}).get(f'{report_id}:{t.payload["record_id"]}')
        allowed={x['key'] for x in r.snapshot.get('selected_requirements',[])}
        if not override or override not in allowed:raise HTTPException(422,'请先在附录D界面确认该条结果记录对应的测评项')
        t.payload={**t.payload,'requirement_key':override,'match_status':'manual_at_retry'}
    # A newly matched record has no approved redacted request yet. Rebuild the
    # pending batch and wait for explicit authorization before queueing work.
    if t.kind=='record' and 'appendix_d' in (r.modules or []) and not (t.output or {}).get('redaction'):
        return _refresh_redaction_preview(db,pid,r,auth[0])
    t.status='queued';t.error='';t.lease='';t.lease_until=0
    r.status='queued';audit(db,auth[0],'task_retried',pid,task_id=tid);db.commit();POOL.submit(drain_tasks)
    return {'ok':True}


def claim_task(db):
    # ponytail: SQLite is a one-process local demo; PostgreSQL also supports row locks.
    for _ in range(8):
        now=time.time();token=uuid.uuid4().hex
        condition=(Task.status=='queued') | ((Task.status=='running') & (Task.lease_until<now))
        stmt=select(Task).where(condition).order_by(Task.created_at).limit(1)
        if not str(engine.url).startswith('sqlite'):stmt=stmt.with_for_update(skip_locked=True)
        t=db.scalar(stmt)
        if not t:return None
        r=db.get(Run,t.run_id)
        if r.cancelled:t.status='cancelled';db.commit();continue
        changed=db.execute(update(Task).where(Task.id==t.id,condition).values(status='running',lease=token,lease_until=now+360,attempts=Task.attempts+1))
        if changed.rowcount!=1:db.rollback();continue
        r.status='running';db.commit();return t.id,token
    return None


def recover_tasks():
    with Session() as db:
        for t in db.scalars(select(Task).where(Task.status=='running',Task.lease_until<time.time())):
            t.status='queued';t.lease='';t.lease_until=0
        db.commit()
    drain_tasks()


def issue_rule_version(run, kind, task=None):
    snapshot=run.snapshot or {}
    if kind=='assets':
        return ' / '.join(str(snapshot[key]) for key in ('mapping_version','algorithm_version') if snapshot.get(key))
    if kind=='risk':
        return ','.join(sorted(snapshot.get('guide_ids',[])))
    ids=snapshot.get('knowledge_ids',[])
    if task:
        key=(task.payload or {}).get('requirement_key')
        requirement=next((item for item in snapshot.get('selected_requirements',[]) if item.get('key')==key),None)
        if requirement:
            ids=requirement.get('knowledge_ids') or [requirement.get('knowledge_id')]
    return ','.join(sorted(str(value) for value in ids if value))


def drain_tasks():
    while True:
        with Session() as db:claimed=claim_task(db)
        if not claimed:return
        tid,token=claimed
        try:
            with Session() as db:
                t=db.get(Task,tid);result=perform_task(db,t)
                r=db.get(Run,t.run_id)
                if r.cancelled:status='cancelled';result={'issues':[],'output':{}}
                else:status=result.get('status','done')
                if t.lease!=token:continue
                for item in result.get('issues',[]):
                    item['machine']={**(item.get('machine') or {}),
                                     'rule_version':issue_rule_version(r,t.kind,t)}
                    db.add(Issue(project_id=t.project_id,run_id=t.run_id,task_id=t.id,**item))
                t.output=result.get('output',{});t.status=status;t.finished_at=time.time();t.lease='';t.lease_until=0
                remaining=db.scalars(select(Task).where(Task.run_id==r.id,Task.id!=tid)).all()
                states=[x.status for x in remaining]+[status]
                if all(x in ('done','blocked','failed','cancelled','awaiting_redaction') for x in states):
                    if any(x == 'awaiting_redaction' for x in states):
                        r.status='awaiting_redaction'
                    else:
                        r.status='cancelled' if r.cancelled else ('partial' if any(x in ('blocked','failed') for x in states) else 'done')
                db.commit()
        except Exception as exc:
            with Session() as db:
                t=db.get(Task,tid)
                if t and t.lease==token:
                    t.status='failed';t.error=str(exc)[:400];t.finished_at=time.time();t.lease='';t.lease_until=0
                    r=db.get(Run,t.run_id);r.status='partial';db.commit()


def perform_task(db,t):
    r=db.get(Run,t.run_id);snap=r.snapshot
    docs={role:db.get(Document,did) for role,did in snap['document_ids'].items()}
    if t.kind=='assets':
        result=compare_consistency({k:{'id':v.id,'status':v.status,'parsed':v.parsed} for k,v in docs.items()},t.payload['scope'])
        return {'output':{**result,'algorithm_version':snap.get('algorithm_version',CONSISTENCY_ALGORITHM_VERSION),
                          'created_at':snap.get('created_at',r.created_at),
                          'categories':snap.get('consistency_categories',CONSISTENCY_CATEGORIES),
                          'document_versions':snap.get('document_versions',{})}}
    if t.kind=='risk':
        guide=[row for kid in snap['guide_ids'] for row in db.get(Knowledge,kid).content.get('requirements',[])]
        result=risk_review(docs['report'].parsed,guide,snap.get('project_config') or {})
        return {'issues':result['issues'],'output':result}
    record=next(x for x in docs['report'].parsed['records'] if x['id']==t.payload['record_id'])
    reqkey=t.payload.get('requirement_key')
    if not reqkey:return {'status':'blocked','output':{'reason':'测评项未能唯一匹配；请人工确认后重试','record':record}}
    req=next((x for x in snap.get('selected_requirements',[]) if x['key']==reqkey),None)
    if not req:return {'status':'blocked','output':{'reason':'核查点版本中未找到测评项','record':record}}
    # Catalog rows without a decision rule or points remain visible as pending
    # local results and must never trigger an external model request.
    if (not isinstance(req.get('decision'),str) or not req.get('decision').strip()
            or not isinstance(req.get('points'),list) or not req['points']):
        classification=classify_appendix(record,req,{},'regular' if r.mode=='lenient' else r.mode)
        evidence=[{'role':'report','source':record.get('source',''),'quote':record.get('text','')}]
        pending_issues=[issue('APP_D','review_pending',item.get('code','analysis_invalid'),
                              item.get('reason',''), 'manual review required', record.get('object',''),
                              evidence, record_id=record.get('id',''), column=item.get('column',''))
                        for item in classification.get('issues',[])]
        return {'status':'blocked','issues':pending_issues,
                'output':{'record':record,'requirement':req,
                          'classification':classification,
                          'reason':'catalog rule is incomplete; manual review required'}}
    profile_snapshot_value=snap.get('model_profile');profile=db.get(UserModelProfile,profile_snapshot_value.get('profile_id')) if profile_snapshot_value else None
    service=db.get(ModelService,profile_snapshot_value.get('service_id')) if profile_snapshot_value else None
    approval=db.scalar(select(RedactionApproval).where(RedactionApproval.run_id==r.id,RedactionApproval.task_id==t.id))
    redaction=(t.output or {}).get('redaction')
    if not profile or not service or not profile_is_current(profile_snapshot_value or {},profile,service):return {'status':'blocked','output':{'reason':'个人模型档案已变更，请重新预览脱敏请求'}}
    if not approval or not redaction: return {'status':'blocked','output':{'reason':'脱敏请求尚未授权'}}
    try:assert_authorized(approval.request_hash,redaction['hash'])
    except ValueError:return {'status':'blocked','output':{'reason':'脱敏请求已变化，请重新预览'}}
    if not str(req.get('decision') or '').strip():
        missing_rule = issue('APP_D', 'description_coverage', '缺少判定规则',
                             '当前核查点没有可用的符合情况判定规则，不能由模型判断为符合。',
                             '补充或确认判定规则后重新审核。', record.get('object', ''),
                             [{'role': 'report', 'source': record.get('source', ''), 'quote': record.get('text', '')},
                              {'role': 'knowledge', 'source': req.get('source', ''), 'quote': req.get('text', '')}],
                             record_id=record.get('id'), column='N', reason_code='missing_decision_rule')
        return {'status': 'blocked', 'issues': [missing_rule],
                'output': {'record': record, 'requirement': req,
                           'reason': '核查点缺少判定规则，待人工核实'}}
    key=Fernet(master_key()).decrypt(profile.key_encrypted.encode()).decode()
    redacted=redaction['payload']
    redacted_record=dict(redacted.get('record') or {})
    redacted_record.setdefault('id', redacted_record.get('record_id') or record.get('id'))
    redacted_req=dict(redacted.get('requirement') or {})
    analysis=evaluate(redacted_record,redacted_req,r.mode,{'base_url':service.base_url,'model':profile.model,'timeout':120},key)
    classified=classify_appendix(redacted_record,redacted_req,analysis,'regular' if r.mode=='lenient' else r.mode)
    batch=db.scalar(select(RedactionBatch).where(RedactionBatch.run_id==r.id))
    reverse={}
    if batch and batch.token_map_encrypted:
        try:
            token_map=json.loads(Fernet(master_key()).decrypt(batch.token_map_encrypted.encode()).decode())
            reverse={token:raw for raw,token in token_map.items()}
        except Exception:
            reverse={}
    def restore(value):
        if isinstance(value,str):
            for token,raw in reverse.items():value=value.replace(token,raw)
            return value
        if isinstance(value,list):return [restore(item) for item in value]
        if isinstance(value,dict):return {key:restore(item) for key,item in value.items()}
        return value
    analysis=restore(analysis)
    classified=restore(classified)
    evidence=[{'role':'report','source':record['source'],'quote':record['text']},
              {'role':'knowledge','source':req['source'],'quote':req['text']},
              {'role':'knowledge','source':req['source']+' · 判定规则','quote':req.get('decision') or '来源未提供判定规则'}]
    issues=[];category_map={'J':'verdict','K':'unrelated','L':'grammar','M':'typo','N':'description_coverage'}
    for column,items in classified['columns'].items():
        for item in items:
            issues.append(issue('APP_D',category_map[column],item.get('code','附录D问题'),item.get('reason',''),item.get('suggestion','人工复核'),record['object'],evidence,record_id=record['id'],column=column,**{k:v for k,v in item.items() if k not in ('reason','suggestion')}))
    for item in classified.get('issues',[]):
        if item.get('column'):
            continue
        issues.append(issue('APP_D','review_pending',item.get('code','analysis_pending'),item.get('reason',''),
                            '由人工核对证据后确认',record['object'],evidence,
                            record_id=record['id'],column='N',pending=True))
    return {'issues':issues,'output':{'record':record,'requirement':req,'analysis':analysis,'redaction':redaction,'classification':classified}}


@app.get('/api/chapters')
def chapters(auth=Depends(current)):
    return [{'code':code,'title':title,'configured':code in ('ALL','CROSS_DOCUMENT','APP_D','CH05','MAJOR_HAZARD')} for code,title in CHAPTERS]


def _issue_rows_for_module(db, rows, module):
    """Apply the module filter against the task snapshot, with legacy fallbacks."""
    if not module:
        return rows
    if module not in MODULES:
        raise HTTPException(422, '模块筛选无效')
    task_ids={row.task_id for row in rows if row.task_id}
    tasks={task.id: task for task in db.scalars(select(Task).where(Task.id.in_(task_ids))).all()} if task_ids else {}
    result=[]
    for row in rows:
        task=tasks.get(row.task_id)
        if task:
            if task.kind == 'record': actual='appendix_d'
            elif task.kind == 'risk': actual='high_risk'
            elif task.kind == 'assets': actual='assets_' + str((task.payload or {}).get('scope', 'full'))
            else: actual=''
        elif row.chapter == 'APP_D': actual='appendix_d'
        elif row.chapter in ('MAJOR_HAZARD', 'CH05'): actual='high_risk'
        elif row.chapter in ('CROSS_DOCUMENT', 'ALL'):
            actual='assets_full'
        else: actual=''
        if actual == module:
            result.append(row)
    return result


def _issue_run(db, pid, run_id):
    if run_id == '':
        return None
    run = db.get(Run, run_id) if run_id is not None else db.scalar(
        select(Run).where(Run.project_id == pid).order_by(Run.created_at.desc(), Run.id.desc()).limit(1))
    if run_id and (not run or run.project_id != pid):
        raise HTTPException(404, '运行不存在')
    return run


def _issue_chapters(row):
    related = (row.machine or {}).get('related_chapters', [])
    return {row.chapter, *(related if isinstance(related, list) else [])}


def _issue_report(db, pid, run):
    if not run:
        return latest_documents(db, pid).get('report')
    report_id = (run.snapshot or {}).get('document_ids', {}).get('report')
    return db.get(Document, report_id) if report_id else None


def _issue_coverage(db, pid, run):
    report = _issue_report(db, pid, run)
    sections = {section.get('code') for section in (report.parsed or {}).get('sections', [])} if report else set()
    tasks = db.scalars(select(Task).where(Task.run_id == run.id)).all() if run else []
    issues = db.scalars(select(Issue).where(Issue.run_id == run.id)).all() if run else []
    modules = {'CROSS_DOCUMENT': {'assets_full', 'assets_sample'}, 'APP_D': {'appendix_d'},
               'CH05': {'high_risk'}, 'MAJOR_HAZARD': {'high_risk'}}
    kinds = {'assets_full': 'assets', 'assets_sample': 'assets', 'appendix_d': 'record', 'high_risk': 'risk'}
    chapters = []
    for code, title in CHAPTERS:
        if code == 'ALL':
            continue
        configured = modules.get(code)
        relevant = [task for task in tasks if configured and any(
            task.kind == kinds[module] and (task.kind != 'assets' or
            (task.payload or {}).get('scope') == module.removeprefix('assets_'))
            for module in configured if module in (run.modules or []))]
        if not configured:
            review_status = '规则待配置'
        elif not run or not configured.intersection(run.modules or []):
            review_status = '未运行'
        elif any(task.status == 'failed' for task in relevant):
            review_status = '部分失败' if any(task.status == 'done' for task in relevant) else '失败'
        elif any(task.status == 'blocked' for task in relevant):
            review_status = '部分失败'
        elif relevant and all(task.status == 'done' for task in relevant):
            review_status = '已完成'
        else:
            review_status = '运行中'
        parsed = (report.parsed or {}) if report else {}
        chapter_failure = any(str(table.get('location','')).startswith(code) and table.get('parse_status') != 'parsed'
                              for table in parsed.get('source_tables', [])) or any(
            code in str(message) for message in parsed.get('diagnostics', []))
        content_status = ('未上传' if not report else '提取失败' if report.status == 'failed' else
                          '提取不完整' if chapter_failure else
                          '已提取' if code in sections or
                          code == 'FULL_TEXT' and any(block.get('region') == 'body' for block in parsed.get('blocks', [])) or
                          code == 'APP_D' and bool(parsed.get('records')) or
                          code == 'CROSS_DOCUMENT' and bool(parsed.get('source_tables')) else '不适用/未出现')
        visible = [item for item in issues if code in _issue_chapters(item)]
        chapters.append({'code': code, 'title': title, 'content_status': content_status,
                         'review_status': review_status, 'issue_count': len(visible),
                         'confirmed_count': sum(item.status == 'confirmed' for item in visible)})
    return {'run_id': run.id if run else None, 'confirmed_count': sum(item.status == 'confirmed' for item in issues),
            'chapters': chapters}


@app.get('/api/projects/{pid}/issues')
def issues(pid:str,chapter:str='ALL',status:str='',run_id:str|None=None,module:str='',auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);run=_issue_run(db,pid,run_id)
    if not run:return []
    stmt=select(Issue).where(Issue.project_id==pid,Issue.run_id==run.id)
    if status:stmt=stmt.where(Issue.status==status)
    rows=db.scalars(stmt.order_by(Issue.created_at.desc())).all()
    if chapter!='ALL':rows=[row for row in rows if chapter in _issue_chapters(row)]
    return [public(x) for x in _issue_rows_for_module(db, rows, module)]


@app.get('/api/projects/{pid}/issues/coverage')
def issue_coverage(pid:str,run_id:str|None=None,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);return _issue_coverage(db,pid,_issue_run(db,pid,run_id))


@app.get('/api/projects/{pid}/issues/chapter-content')
def issue_chapter_content(pid:str,chapter:str,run_id:str|None=None,
                          auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    if chapter not in {code for code,_ in CHAPTERS}:
        raise HTTPException(422,'章节无效')
    run=_issue_run(db,pid,run_id)
    report=_issue_report(db,pid,run)
    coverage=_issue_coverage(db,pid,run)
    status=next((item for item in coverage['chapters'] if item['code']==chapter),None)
    if not report:
        return {'run_id':run.id if run else None,'chapter':chapter,'document':None,
                'sections':[],'blocks':[],'content_status':'未上传',
                'review_status':status['review_status'] if status else '未运行'}
    parsed=report.parsed or {}
    sections=[item for item in parsed.get('sections',[]) if chapter=='FULL_TEXT' or item.get('code')==chapter]
    section_ids={item.get('id') for item in sections}
    tables={item.get('block_id'):item for item in parsed.get('source_tables',[])
            if not item.get('parent_table_id')}
    blocks=[]
    for item in parsed.get('blocks',[]):
        if item.get('region')!='body' or item.get('type')=='toc':
            continue
        if chapter!='FULL_TEXT' and item.get('section_id') not in section_ids:
            continue
        block=dict(item)
        if item.get('type')=='table' and item.get('id') in tables:
            block['table']=tables[item['id']]
            block['nested_tables_data']=[table for table in parsed.get('source_tables',[])
                                         if table.get('block_id')==item['id'] and table.get('parent_table_id')]
        blocks.append(block)
    return {'run_id':run.id if run else None,'chapter':chapter,
            'document':{'id':report.id,'filename':report.filename,'version':report.version,
                        'sha256':report.sha256,'status':report.status},
            'sections':sections,'blocks':blocks,
            'content_status':status['content_status'] if status else report.status,
            'review_status':status['review_status'] if status else '规则待配置'}


@app.post('/api/projects/{pid}/issues/batch')
def review_issues_batch(pid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);run_id=body.get('run_id')
    if not run_id:raise HTTPException(422,'缺少运行ID')
    run=_issue_run(db,pid,run_id)
    items=body.get('items');status=body.get('status')
    note=body.get('note','')
    note=note[:4000] if isinstance(note,str) else ''
    if not isinstance(items,list) or not items or len(items)>500 or status not in ('pending','confirmed','rejected','needs_evidence'):
        raise HTTPException(422,'批量复核参数无效')
    if status in ('rejected','needs_evidence') and not note.strip():
        raise HTTPException(422,'请填写复核理由')
    ids=[item.get('id') for item in items if isinstance(item,dict)]
    if len(ids)!=len(items) or any(not isinstance(iid,str) or not iid for iid in ids) or len(set(ids))!=len(ids):
        raise HTTPException(422,'问题ID重复或无效')
    rows={row.id:row for row in db.scalars(select(Issue).where(Issue.id.in_(ids))).all()}
    conflicts=[iid for iid,item in zip(ids,items) if iid not in rows or rows[iid].project_id!=pid
               or rows[iid].run_id!=run.id or rows[iid].version!=item.get('version')]
    if conflicts:raise HTTPException(409,{'conflicts':conflicts})
    before={iid:(rows[iid].status,rows[iid].version) for iid in ids}
    for iid in ids:
        result=db.execute(update(Issue).where(Issue.id==iid,Issue.version==before[iid][1])
                          .values(status=status,note=note,version=before[iid][1]+1))
        if result.rowcount!=1:
            db.rollback();raise HTTPException(409,{'conflicts':[iid]})
    audit(db,auth[0],'issues_batch_reviewed',pid,run_id=run.id,issue_ids=ids,
          before={iid:{'status':s,'version':v} for iid,(s,v) in before.items()},status=status,note=note)
    db.commit()
    return {'items':[public(db.get(Issue,iid)) for iid in ids]}


@app.patch('/api/projects/{pid}/issues/{iid}')
def review_issue(pid:str,iid:str,body:dict,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0]);row=db.get(Issue,iid)
    if not row or row.project_id!=pid:raise HTTPException(404,'问题不存在')
    if row.version!=body.get('version'):raise HTTPException(409,'问题已被其他审核员修改，请刷新')
    status=body.get('status',row.status)
    if status not in ('pending','confirmed','rejected','needs_evidence','resolved'):raise HTTPException(422,'复核状态无效')
    note=body.get('note',row.note)
    note=note[:4000] if isinstance(note,str) else ''
    if status in ('rejected','needs_evidence') and not note.strip():raise HTTPException(422,'请填写复核理由')
    before={'status':row.status,'version':row.version}
    changed=db.execute(update(Issue).where(Issue.id==iid,Issue.version==row.version)
                       .values(status=status,note=note,version=row.version+1))
    if changed.rowcount!=1:
        db.rollback();raise HTTPException(409,'问题已被其他审核员修改，请刷新')
    audit(db,auth[0],'issue_reviewed',pid,issue_id=iid,before=before,status=status,version=before['version']+1,note=note)
    db.commit();db.refresh(row);return public(row)


def _risk_task(db, pid, rid):
    run = db.get(Run, rid)
    if not run or run.project_id != pid:
        raise HTTPException(404, '运行不存在')
    task = db.scalar(select(Task).where(Task.run_id == rid, Task.kind == 'risk'))
    if not task or task.status not in ('done', 'partial'):
        raise HTTPException(409, '高风险核查尚未完成')
    return task


def _risk_output_for_task(db, task):
    output = json.loads(json.dumps(task.output or {}, ensure_ascii=False))
    if 'legacy_rows' not in output:
        output['legacy_rows'] = legacy_risk_rows(output.get('source_rows', []), output.get('guide_rows', []))
    if output.get('guide_rows') and not any(row.get('requirement') for row in output['guide_rows']):
        output['legacy_note'] = '当前运行使用的旧指引快照缺少“要求项/标准要求”；重新导入指引并发起新运行后可生成旧版13列汇总。'
    decisions = {row.row_id: _risk_decision_public(row) for row in db.scalars(
        select(ReviewDecision).where(ReviewDecision.run_id == task.run_id))}
    for group in ('source_rows', 'candidates', 'links'):
        for row in output.get(group, []):
            key = {'source_rows': 'source:', 'candidates': 'candidate:', 'links': 'link:'}[group]
            row['review'] = decisions.get(key + str(row.get('row_id') or row.get('link_id')))
    for row in output['legacy_rows']:
        row['review'] = decisions.get('candidate:' + str(row.get('source_row_id')))
    output['run_id'] = task.run_id
    return output


def _risk_decision_public(row):
    value = public(row)
    try:
        detail = json.loads(row.note or '{}')
    except (TypeError, ValueError):
        detail = {'reason': row.note or ''}
    if not isinstance(detail, dict):
        detail = {'reason': row.note or ''}
    value['conclusion'] = str(detail.get('conclusion', ''))
    value['reason'] = str(detail.get('reason', ''))
    return value


@app.get('/api/projects/{pid}/runs/{rid}/risk')
def run_risk_output(pid: str, rid: str, auth=Depends(current), db: Session=Depends(db_session)):
    project_access(db, pid, auth[0])
    return _risk_output_for_task(db, _risk_task(db, pid, rid))


@app.patch('/api/projects/{pid}/runs/{rid}/risk/{row_id}')
def review_risk_row(pid: str, rid: str, row_id: str, body: dict, auth=Depends(current), db: Session=Depends(db_session)):
    project_access(db, pid, auth[0]); task = _risk_task(db, pid, rid); output = task.output or {}
    valid = {f"source:{x.get('row_id')}" for x in output.get('source_rows', [])}
    valid |= {f"candidate:{x.get('row_id')}" for x in output.get('candidates', [])}
    valid |= {f"link:{x.get('link_id')}" for x in output.get('links', [])}
    key = row_id if row_id in valid else next((x for x in valid if x.endswith(':' + row_id)), None)
    if not key: raise HTTPException(404, '高风险核查行不存在')
    status = body.get('status', 'pending')
    if status not in ('pending', 'confirmed', 'rejected', 'needs_evidence', 'not_applicable', 'linked'):
        raise HTTPException(422, '复核状态无效')
    existing = db.scalar(select(ReviewDecision).where(ReviewDecision.run_id == rid, ReviewDecision.row_id == key))
    expected = existing.version if existing else 0
    if body.get('version') != expected: raise HTTPException(409, '复核已被修改，请刷新后重试')
    detail = {'conclusion': str(body.get('conclusion', ''))[:2000], 'reason': str(body.get('reason', body.get('note', '')))[:4000]}
    if existing: existing.version += 1; decision = existing
    else: decision = ReviewDecision(run_id=rid, row_id=key, reviewer_id=auth[0].id, version=1); db.add(decision)
    decision.status = status; decision.note = json.dumps(detail, ensure_ascii=False); decision.reviewer_id = auth[0].id; decision.updated_at = time.time()
    audit(db, auth[0], 'risk_reviewed', pid, run_id=rid, row_id=key, status=status); db.commit()
    return {**public(decision), **detail}


@app.get('/api/projects/{pid}/runs/{rid}/risk/export')
def export_risk_run(pid: str, rid: str, auth=Depends(current), db: Session=Depends(db_session)):
    project_access(db, pid, auth[0]); task = _risk_task(db, pid, rid); result = _risk_output_for_task(db, task)
    book = Workbook(); legacy = book.active; legacy.title = '高风险核查汇总'
    append_text_row(legacy, LEGACY_RISK_COLUMNS)
    for row in result.get('legacy_rows', []):
        append_text_row(legacy, [row.get('values', {}).get(name, '') for name in LEGACY_RISK_COLUMNS])
    legacy.freeze_panes = 'A2'
    legacy.auto_filter.ref = legacy.dimensions
    for column, width in zip('ABCDEFGHIJKLM', (8, 12, 32, 32, 40, 26, 18, 28, 14, 18, 42, 42, 42)):
        legacy.column_dimensions[column].width = width
    for cell in legacy[1]:
        cell.fill = PatternFill('solid', fgColor='DDECF0')
        cell.font = Font(bold=True, color='244857')
    for row in legacy.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        major = str(row[9].value or '')
        if '重大风险项' in major:
            row[9].fill = PatternFill('solid', fgColor='FFC0CB')
        elif '高风险项' in major:
            row[9].fill = PatternFill('solid', fgColor='FFE699')
    # The desktop export merges adjacent equal guide text in K-M.
    for col in (11, 12, 13):
        start = 2
        while start <= legacy.max_row:
            value = str(legacy.cell(start, col).value or '').strip()
            end = start
            while end < legacy.max_row and str(legacy.cell(end + 1, col).value or '').strip() == value:
                end += 1
            if end > start:
                legacy.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            start = end + 1
    sheet = book.create_sheet('风险来源行')
    append_text_row(sheet, ['运行ID', '来源行ID', '章节', '来源位置', '对象', '描述', '风险分析', '风险等级', '原始值', '人工状态', '人工结论', '复核理由'])
    for row in result.get('source_rows', []):
        review = row.get('review') or {}
        append_text_row(sheet, [rid, row.get('row_id'), row.get('chapter'), row.get('table_location'), row.get('object'), row.get('description'), row.get('analysis'), row.get('grade'), json.dumps(row.get('raw_values', {}), ensure_ascii=False), review.get('status', 'pending'), review.get('conclusion', ''), review.get('reason', '')])

    source_detail = book.create_sheet('\u6765\u6e90\u884c\u660e\u7ec6')
    append_text_row(source_detail, ['\u6765\u6e90\u884cID', '\u7ae0\u8282', '\u8868\u683c\u4f4d\u7f6e', '\u884c\u53f7', '\u8868\u7c7b\u578b', '\u89e3\u6790\u72b6\u6001', '\u5bf9\u8c61', '\u95ee\u9898\u63cf\u8ff0', '\u98ce\u9669\u5206\u6790', '\u98ce\u9669\u7b49\u7ea7', '\u539f\u59cb\u5b57\u6bb5\u503c', '\u590d\u6838\u72b6\u6001', '\u4eba\u5de5\u7ed3\u8bba', '\u590d\u6838\u7406\u7531'])
    for row in result.get('source_rows', []):
        review = row.get('review') or {}
        append_text_row(source_detail, [row.get('row_id'), row.get('chapter'), row.get('table_location'), row.get('row_number'), row.get('kind'), row.get('parse_status'), row.get('object'), row.get('description'), row.get('analysis'), row.get('grade'), json.dumps(row.get('raw_values', {}), ensure_ascii=False), review.get('status', 'pending'), review.get('conclusion', ''), review.get('reason', '')])
    candidates = book.create_sheet('\u5019\u9009\u5173\u8054')
    append_text_row(candidates, ['source_row_id', 'guide_key', 'guide_source', 'scenario', 'score', 'applicability', 'status', 'machine_conclusion', 'review_status', 'human_conclusion', 'review_reason'])
    guide_rows = {str(row.get('key', row.get('id', ''))): row for row in result.get('guide_rows', [])}
    candidates_by_row = {row.get('row_id'): row for row in result.get('candidates', [])}
    for link in result.get('links', []):
        review = link.get('review') or candidates_by_row.get(link.get('source_row_id'), {}).get('review') or {}
        guide_key = str(link.get('guide_key') or link.get('candidate_key') or '')
        guide = guide_rows.get(guide_key, {})
        append_text_row(candidates, [link.get('source_row_id'), guide_key, link.get('guide_source') or guide.get('source', ''), guide.get('scenario') or guide.get('description') or guide.get('text', ''), link.get('score'), link.get('applicability'), link.get('status'), link.get('machine_conclusion'), review.get('status', 'pending'), review.get('conclusion', ''), review.get('reason', '')])
    for candidate in result.get('candidates', []):
        if candidate.get('matches'):
            continue
        review = candidate.get('review') or {}
        append_text_row(candidates, [candidate.get('row_id'), '', '', '', '', '', candidate.get('status', 'unmatched'), '', review.get('status', 'pending'), review.get('conclusion', ''), review.get('reason', '')])

    unmatched = book.create_sheet('\u672a\u5339\u914d\u6765\u6e90')
    append_text_row(unmatched, ['source_row_id', 'source', 'reason'])
    for row in result.get('unmatched', []):
        append_text_row(unmatched, [row.get('source_row_id') or row.get('row_id'), row.get('source', ''), row.get('reason', '')])

    conflicts = book.create_sheet('\u51b2\u7a81\u4e0e\u9002\u7528\u6027')
    append_text_row(conflicts, ['type', 'row_id', 'expected', 'actual', 'source', 'guide_key', 'applicability'])
    for row in result.get('conflicts', []):
        append_text_row(conflicts, [row.get('type'), row.get('row_id'), row.get('expected'), row.get('actual'), row.get('source')])
    for key, status in (result.get('applicability') or {}).items():
        guide = guide_rows.get(str(key), {})
        append_text_row(conflicts, ['applicability', '', '', '', guide.get('source', ''), key, status])

    evidence = book.create_sheet('\u56db\u65b9\u6838\u5bf9\u8bc1\u636e')
    append_text_row(evidence, ['\u8bc1\u636e\u7c7b\u578b', '\u6765\u6e90\u884cID', '\u6307\u5f15ID', '\u6765\u6e90\u4f4d\u7f6e', '\u5185\u5bb9'])
    for row in result.get('source_rows', []):
        append_text_row(evidence, ['report', row.get('row_id'), '', row.get('table_location') or row.get('source', ''), json.dumps(row.get('raw_values', {}), ensure_ascii=False)])
    for row in result.get('guide_rows', []):
        append_text_row(evidence, ['guide', '', row.get('key') or row.get('id'), row.get('source', ''), json.dumps(row, ensure_ascii=False)])
    for row in result.get('links', []):
        guide_key = row.get('guide_key') or row.get('candidate_key')
        guide = guide_rows.get(str(guide_key), {})
        append_text_row(evidence, ['link', row.get('source_row_id'), guide_key, guide.get('source', row.get('guide_source', '')), json.dumps(row, ensure_ascii=False)])
    for row in result.get('conflicts', []):
        append_text_row(evidence, ['conflict', row.get('row_id'), '', row.get('source', ''), json.dumps(row, ensure_ascii=False)])

    stream = io.BytesIO(); book.save(stream); stream.seek(0); audit(db, auth[0], 'risk_exported', pid, run_id=rid); db.commit()
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="risk-{rid}.xlsx"'})


@app.get('/api/projects/{pid}/risk')
def risk_output(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    tasks=db.scalars(select(Task).where(Task.project_id==pid,Task.kind=='risk',Task.status=='done').order_by(Task.finished_at.desc())).all()
    return _risk_output_for_task(db, tasks[0]) if tasks else {'candidates':[],'linked_rows':[],'note':'尚未运行高风险核查'}


@app.get('/api/projects/{pid}/export')
def export_issues(pid:str,format:str='xlsx',run_id:str|None=None,module:str='',chapter:str='ALL',status:str='confirmed',auth=Depends(current),db=Depends(db_session)):
    p=project_access(db,pid,auth[0]);run=_issue_run(db,pid,run_id)
    if status not in ('confirmed','all'):raise HTTPException(422,'导出状态无效')
    stmt=select(Issue).where(Issue.project_id==pid,Issue.run_id==run.id) if run else None
    if stmt is not None and status=='confirmed':stmt=stmt.where(Issue.status=='confirmed')
    rows=db.scalars(stmt.order_by(Issue.chapter,Issue.created_at)).all() if stmt is not None else []
    if chapter!='ALL':rows=[row for row in rows if chapter in _issue_chapters(row)]
    rows=_issue_rows_for_module(db, rows, module)
    coverage=_issue_coverage(db,pid,run)
    headers=['章节','问题类别','问题标题','问题说明','修改建议','关联对象','原文位置','复核意见','关联章节','规则版本','复核状态']
    values=[[x.chapter,x.category,x.title,x.description,x.suggestion,x.object_name,
             '；'.join(e.get('source','') for e in x.evidence),x.note,
             '、'.join(sorted(_issue_chapters(x)-{x.chapter})),(x.machine or {}).get('rule_version',''),x.status] for x in rows]
    summary=f'{len(rows)}条已确认问题' if status=='confirmed' else f'{len(rows)}条候选问题'
    if not rows and status=='confirmed':summary+='；不能据此判定报告无问题'
    generated=time.strftime('%Y-%m-%d %H:%M:%S',time.localtime())
    if format=='xlsx':
        book=Workbook();sheet=book.active;sheet.title='已确认问题' if status=='confirmed' else '全部问题';append_text_row(sheet,headers)
        # Imported report text must stay text when opened in Excel.
        for row in values:append_text_row(sheet,row)
        for col,width in {'A':18,'B':22,'C':32,'D':60,'E':60,'F':28,'G':46,'H':40}.items():sheet.column_dimensions[col].width=width
        sheet.column_dimensions['I'].width=28
        sheet.column_dimensions['J'].width=22
        sheet.column_dimensions['K'].width=16
        cover=book.create_sheet('审核覆盖情况')
        for row in [('项目',p.name),('运行ID',run.id if run else '无'),('导出范围',chapter),('模块',module or '全部'),('状态',status),('生成时间',generated),('说明',summary),
                    ('章节','内容状态','审核状态','候选问题数','已确认数')]:append_text_row(cover,row)
        for item in coverage['chapters']:
            cover.append([item['title'],item['content_status'],item['review_status'],item['issue_count'],item['confirmed_count']])
        stream=io.BytesIO();book.save(stream);media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet';suffix='xlsx'
    elif format=='docx':
        book=WordDocument();book.add_heading('等保测评报告审核问题汇总',0)
        book.add_paragraph(f'项目：{p.name}；运行ID：{run.id if run else "无"}；导出范围：{chapter}；模块：{module or "全部"}；状态：{status}；生成时间：{generated}。')
        book.add_paragraph(summary)
        current_chapter=None
        for i,row in enumerate(values,1):
            if row[0]!=current_chapter:
                current_chapter=row[0]
                book.add_heading(current_chapter,level=1)
            book.add_heading(f'{i}. {row[2]}',level=2)
            for name,val in zip(headers,row):book.add_paragraph(f'{name}：{val}')
        book.add_heading('审核覆盖情况',level=1)
        for item in coverage['chapters']:
            book.add_paragraph(f'{item["title"]}：内容{item["content_status"]}；审核{item["review_status"]}；候选{item["issue_count"]}条；已确认{item["confirmed_count"]}条。')
        stream=io.BytesIO();book.save(stream);media='application/vnd.openxmlformats-officedocument.wordprocessingml.document';suffix='docx'
    else:raise HTTPException(422,'仅支持xlsx或docx')
    stream.seek(0);audit(db,auth[0],'issues_exported',pid,run_id=run.id if run else None,format=format,module=module,chapter=chapter,status=status,count=len(rows));db.commit()
    return StreamingResponse(stream,media_type=media,headers={'Content-Disposition':f'attachment; filename="review-issues.{suffix}"'})


@app.get('/api/projects/{pid}/audit')
def audit_log(pid:str,auth=Depends(current),db=Depends(db_session)):
    project_access(db,pid,auth[0])
    return [public(x) for x in db.scalars(select(Audit).where(Audit.project_id==pid).order_by(Audit.created_at.desc()).limit(200))]


DIST = Path(__file__).resolve().parent.parent / 'frontend' / 'dist'
if DIST.exists():
    app.mount('/assets',StaticFiles(directory=DIST/'assets'),name='assets')


@app.get('/')
@app.get('/{path:path}')
def frontend(path:str=''):
    if path.startswith('api/'):raise HTTPException(404,'接口不存在')
    if not DIST.exists():return {'message':'前端尚未构建，请在 frontend 执行 npm run build'}
    return FileResponse(DIST/'index.html')

