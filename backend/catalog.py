"""Read user-maintained workbooks. Never edit source workbooks or invent missing rules."""
import hashlib
import re
import unicodedata
from collections import Counter
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook


def norm(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value or ''))).casefold()


def safe_zip(data, required):
    try:
        with ZipFile(BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 15000 or sum(i.file_size for i in infos) > 300 * 1024 * 1024:
                raise ValueError('文档解压后过大')
            if required not in z.namelist():
                raise ValueError('文件内容与格式不一致')
    except BadZipFile as exc:
        raise ValueError('不是有效的 Office 文档') from exc


def split_points(text):
    text = unicodedata.normalize('NFKC', str(text or '')).strip()
    matches = list(re.finditer(r'(?:^|[\n;；。])\s*(\d+)\s*[)、.]\s*', text))
    if not matches:
        return [text] if text else []
    if [int(m[1]) for m in matches] != list(range(1, len(matches) + 1)):
        return [text]
    return [text[m.end():matches[i+1].start() if i+1 < len(matches) else len(text)].strip() for i, m in enumerate(matches)]


def import_workbook(data, filename):
    safe_zip(data, 'xl/workbook.xml')
    wb = load_workbook(BytesIO(data), data_only=False)
    if '附录A' in filename:
        profiles = []
        for sheet in wb:
            if 'A.1' in sheet.title:
                for row in sheet.iter_rows(values_only=True):
                    text = ' '.join(str(v) for v in row if v)
                    for s, a in re.findall(r'S(\d)A(\d)', text):
                        profiles.append({'s': int(s), 'a': int(a), 'g': max(int(s), int(a))})
        return {'family': 'mapping', 'level': None, 'profile': '', 'requirements': [], 'profiles': profiles, 'warnings': [], 'sheets': wb.sheetnames}
    if '高风险' in filename:
        return import_guide(wb)
    families = [('安全通用', 'general'), ('云计算', 'cloud'), ('移动互联', 'mobile'), ('物联网', 'iot'),
                ('工业控制', 'ics'), ('电力监控系统', 'power_monitoring'),
                ('电力管理信息系统', 'power_management'), ('大数据', 'bigdata')]
    family = next((v for k, v in families if k in filename), None)
    if not family:
        raise ValueError('无法识别依据类型，请保留原文件的通用/扩展类型名称')
    level = next((v for k, v in [('二级', 2), ('三级', 3), ('四级', 4)] if k in filename), None)
    profile = re.search(r'S\dA\dG\d', filename)
    requirements, warnings = [], []
    for sh in wb:
        seven = 'SAG' in str(sh.cell(1, 3).value).upper()
        if not seven and '测评项' not in str(sh.cell(1, 3).value):
            warnings.append(f'{sh.title}：表头未识别，未导入'); continue
        merged = {}
        col = 4 if seven else 3
        group_rows = set()
        for area in sh.merged_cells.ranges:
            if area.min_row == area.max_row and area.min_col == 1 and area.max_col >= col:
                group_rows.add(area.min_row)
            for r in range(area.min_row, area.max_row + 1):
                for c in range(area.min_col, area.max_col + 1):
                    merged[(r,c)] = sh.cell(area.min_row, area.min_col).value
        for r in range(2, sh.max_row + 1):
            if r in group_rows:
                continue
            text = merged.get((r, col), sh.cell(r, col).value)
            if not text:
                continue
            sag = str(sh.cell(r, 3).value or '').strip().upper() if seven else ''
            raw = str(merged.get((r, col + 2), sh.cell(r, col + 2).value) or '')
            points = split_points(raw)
            count = sh.cell(r, col + 1).value
            rule = str(merged.get((r, 7), sh.cell(r, 7).value) or '') if seven else ''
            notes = []
            if seven and sag not in ('S', 'A', 'G'): notes.append('S/A/G标识缺失或非法')
            if isinstance(count, (int, float)) and count != len(points): notes.append(f'原点数{count}与解析点数{len(points)}不同')
            if not rule: notes.append('缺少判定规则，仅可检查描述覆盖')
            if len(set(norm(x) for x in points)) != len(points): notes.append('存在重复核查点')
            if any(x in str(text) + raw for x in ('盯水渗洲','配摺','中国极内','抗策')): notes.append('疑似来源识别错字，请核实')
            key = hashlib.sha256(f'{filename}|{sh.title}|{r}'.encode()).hexdigest()[:24]
            source = f'{filename} · {sh.title} · {r}行'
            requirements.append({'key': key, 'domain': sh.title, 'control': str(merged.get((r,2), sh.cell(r,2).value) or ''), 'sag': sag, 'text': str(text), 'points': [{'id': f'{key}:{i+1}', 'text': t, 'sources': [source]} for i,t in enumerate(points)], 'decision': rule, 'declared_count': count, 'source': source, 'sources': [source], 'source_sheet': sh.title, 'source_row': r, 'warnings': notes, 'available': bool(points) and (not seven or sag in ('S','A','G'))})
            warnings.extend(f'{sh.title}!{r}: {n}' for n in notes)
    if not requirements: raise ValueError('未找到有效测评项')
    return {'family': family, 'level': level, 'profile': profile[0] if profile else '', 'requirements': requirements, 'warnings': warnings, 'sheets': wb.sheetnames}


def import_guide(wb):
    sh = wb['扩充文档'] if '扩充文档' in wb.sheetnames else wb.worksheets[0]
    merged = {}
    for area in sh.merged_cells.ranges:
        for r in range(area.min_row, area.max_row+1):
            for c in range(area.min_col, area.max_col+1): merged[(r,c)] = sh.cell(area.min_row,area.min_col).value
    rows=[]
    # Headers are located from labels, not assumed fixed column positions.
    header_row, columns = None, {}
    aliases={'clause':['条款号','条款编号'], 'requirement':['测评指标','测评要求','要求项','标准要求'], 'scope':['适用范围'], 'scenario':['场景','问题描述'], 'mitigation':['缓解'], 'evaluation':['风险评价'], 'major':['重大']}
    for r in range(1,min(sh.max_row,12)+1):
        found={}
        for c in range(1,sh.max_column+1):
            value=str(sh.cell(r,c).value or '')
            for field,words in aliases.items():
                if any(w in value for w in words) and field not in found: found[field]=c
        if 'scenario' in found and len(found)>=3: header_row,columns=r,found; break
    if not header_row: raise ValueError('高风险指引表头无法识别')
    previous={}
    for r in range(header_row+1,sh.max_row+1):
        raw={k:str(merged.get((r,c),sh.cell(r,c).value) or '') for k,c in columns.items()}
        if not any(raw.values()): continue
        vals={k:raw[k] or previous.get(k,'') for k in columns}
        previous=vals
        if not vals.get('scenario') or not vals.get('clause') or not re.search(r'\d',vals['clause']): continue
        if len(vals.get('scenario',''))<10: continue
        rows.append({'key':f'guide:{sh.title}:{r}', **vals, 'source':f'{sh.title} · {r}行'})
    return {'family':'high_risk','level':None,'profile':'','requirements':rows,'warnings':[],'sheets':wb.sheetnames}


def select_requirements(catalogs, config):
    s,a,g=[int(config.get(k,3)) for k in ('s','a','g')]
    extensions=set(config.get('extensions',[]))
    power=config.get('power_category')
    choices={'power_monitoring','power_management'}
    if power and power not in choices:
        raise ValueError('电力扩展类别无效')
    if 'power' in extensions or len(extensions & choices)>1 or (power and (extensions & choices) - {power}):
        raise ValueError('电力监控系统与电力管理信息系统必须二选一')
    if power: extensions.add(power)
    selected=[]; by_item={}
    for cat in catalogs:
        if cat['family'] not in {'general', *extensions}: continue
        for req in cat['content']['requirements']:
            if not req.get('available'): continue
            if cat['profile']:
                ok=cat['profile']==f'S{s}A{a}G{g}'
            else:
                ok=cat['level']=={'S':s,'A':a,'G':g}.get(req['sag'])
            if not ok: continue
            item={**req, 'points':[{**point, 'sources':list(point.get('sources') or [req['source']])} for point in req['points']],
                  'sources':list(req.get('sources') or [req['source']]), 'family':cat['family'],
                  'knowledge_id':cat['id'], 'knowledge_ids':[cat['id']], 'source_level':cat['level']}
            identity=(norm(req['domain']),norm(req['control']),norm(req['text']),req['sag'])
            siblings=by_item.setdefault(identity,[])
            point_names={norm(point['text']) for point in item['points']}
            for sibling in siblings:
                sibling_names={norm(point['text']) for point in sibling['points']}
                if point_names & sibling_names and norm(sibling['decision'])!=norm(item['decision']):
                    raise ValueError(f'判定规则冲突：{sibling["source"]} / {item["source"]}')
            previous=next((x for x in siblings if norm(x['decision'])==norm(item['decision'])),None)
            if previous is None:
                # The same measurement text may occur in different source rows
                # with distinct points and distinct rules. Keep those rows separate.
                siblings.append(item);selected.append(item);continue
            # General and the chosen power supplement may repeat the same point.
            # Preserve both exact source locations; a rule disagreement needs review.
            previous['sources'].extend(item['sources'])
            previous['knowledge_ids'].append(cat['id'])
            points={norm(point['text']):point for point in previous['points']}
            for point in item['points']:
                existing=points.get(norm(point['text']))
                if existing: existing['sources'].extend(point['sources'])
                else: previous['points'].append(point);points[norm(point['text'])]=point
    return selected


def selection_summary(reqs):
    return {'count':len(reqs),'by_attribute':dict(Counter(x['sag'] or '旧格式' for x in reqs)), 'by_family':dict(Counter(x['family'] for x in reqs)), 'without_decision':sum(not x['decision'] for x in reqs)}

